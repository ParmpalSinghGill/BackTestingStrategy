"""
Best matrix variant + 2nd attempt, full 1:2 and 1:3 portfolio returns.

Rules:
  C1 green, open below support
  A1: C2 close > C1 high → enter C3 at C3 open
  SL = lowest low of sweep-to-C1 × 0.99
  If attempt 1 hits SL before 1:2, last SL becomes next liquidity; re-enter once (attempt 2)
  Full exit at 1:2 or 1:3 (separate runs)
  Size: min(floor(risk_cap / risk_share), floor(cash / price)); skip if 1 share risks > cap
"""

from __future__ import annotations

import sys
from collections import defaultdict
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from src.analysis.indian_brokerage_calculator import calculate_indian_trade_charges
from src.backtest_engine.backtest_support_liquidity_strategy import (
    INDEX_CLASSIFIER,
    get_all_stock_supports,
)
from swing_strategy.run_c1_entry_sl_matrix import (
    MAX_POST_SWEEP,
    START,
    _c1_match,
    _entries_for_c1,
    _sl_price,
)
from swing_strategy.tiered_liquidity_strategy_engine import (
    DATA_DAILY_DIR,
    NIFTY_RANK,
    TF_RANK,
    _load_daily,
    _resolve_effective_liquidity,
)
from swing_strategy.visualizer import generate_all_visualizations

OUT_BASE = BASE_DIR / "Reports" / "Retry2_RR_Returns"
PLOT_BASE = BASE_DIR / "Plots" / "Retry2_RR_Returns"
CK, EK, SK = "OPEN_BELOW", "A1", "SWEEP_x99"
MAX_RETRY_WAIT = 252
GAP = 0.001
RULE = "OPEN_BELOW + A1 + sweep×0.99 SL + 2nd attempt if SL before 1:2"

SCENARIOS = [
    {"rr": 2.0, "capital": 50_000.0, "risk": 500.0, "name": "RR1to2_Cap50k_Risk500"},
    {"rr": 3.0, "capital": 50_000.0, "risk": 500.0, "name": "RR1to3_Cap50k_Risk500"},
    {"rr": 2.0, "capital": 100_000.0, "risk": 500.0, "name": "RR1to2_Cap100k_Risk500"},
    {"rr": 3.0, "capital": 100_000.0, "risk": 500.0, "name": "RR1to3_Cap100k_Risk500"},
    {"rr": 2.0, "capital": 50_000.0, "risk": 1_000.0, "name": "RR1to2_Cap50k_Risk1k"},
    {"rr": 3.0, "capital": 50_000.0, "risk": 1_000.0, "name": "RR1to3_Cap50k_Risk1k"},
]


def _try_setup(sweep_idx, support, opens, highs, lows, closes, n):
    end = min(n - 1, sweep_idx + MAX_POST_SWEEP)
    for c1 in range(sweep_idx, end):
        if not _c1_match(CK, float(opens[c1]), float(highs[c1]), float(closes[c1]), support):
            continue
        ents = _entries_for_c1(c1, opens, highs, lows, closes, n)
        if EK not in ents:
            continue
        eidx, entry = ents[EK]
        if entry <= 0:
            continue
        sl = _sl_price(SK, float(lows[c1]), float(min(lows[sweep_idx : c1 + 1])))
        if entry - sl <= 0.05:
            continue
        return int(eidx), float(entry), float(sl), int(c1)
    return None


def _rr_exit(entry_idx, entry, sl, rr, opens, highs, lows, n):
    risk = entry - sl
    tp = round(entry + rr * risk, 2)
    mfe = 0.0
    for m in range(entry_idx, n):
        o, h, l = float(opens[m]), float(highs[m]), float(lows[m])
        if m > entry_idx and o < sl:
            return round(o * (1.0 - GAP), 2), m, mfe, True, "SL"
        if m > entry_idx and o > tp:
            return round(o * (1.0 - GAP), 2), m, max(mfe, (h - entry) / risk), False, "TP"
        mfe = max(mfe, (h - entry) / risk)
        if l <= sl:
            return sl, m, mfe, True, "SL"
        if h >= tp:
            return tp, m, mfe, False, "TP"
    last = n - 1
    return round(float(opens[last]), 2), last, mfe, False, "EOD"


def _next_sweep(start, level, lows, n):
    end = min(n, start + MAX_RETRY_WAIT)
    for j in range(start, end):
        if float(lows[j]) < level:
            return j
    return None


def _pack(symbol, tf, nifty, support, eidx, xidx, entry, sl, exit_p, reason, attempt, date_s, rr):
    return {
        "Ticker": symbol,
        "Liquidity_Type": tf,
        "TF_Rank": TF_RANK.get(tf, 1),
        "Nifty_Rank": nifty,
        "Support_Price": round(support, 2),
        "Entry_Date": date_s.iloc[eidx].strftime("%Y-%m-%d"),
        "Exit_Date": date_s.iloc[xidx].strftime("%Y-%m-%d"),
        "Entry_Price": round(entry, 2),
        "SL_Price": round(sl, 2),
        "Exit_Price": round(exit_p, 2),
        "Attempt": attempt,
        "Outcome": "Success" if reason == "TP" else "Failure",
        "Target_RR_Mode": f"1:{int(rr)}",
    }


def scan_ticker(symbol: str) -> dict[float, list[dict]]:
    df = _load_daily(symbol)
    empty = {2.0: [], 3.0: []}
    if df is None or len(df) < 120:
        return empty
    nifty = NIFTY_RANK.get(INDEX_CLASSIFIER.classify(symbol), 1)
    all_supports = get_all_stock_supports(df.set_index("Date"))
    dates_np = df["Date"].to_numpy()
    date_s = pd.to_datetime(df["Date"])
    opens = df["Open"].to_numpy(float)
    highs = df["High"].to_numpy(float)
    lows = df["Low"].to_numpy(float)
    closes = df["Close"].to_numpy(float)
    n = len(df)

    sup_by_date: dict = {}
    for s in all_supports:
        sup_by_date.setdefault(s["formed_date"], []).append(s)

    active: list[dict] = []
    raw: dict[float, list[dict]] = {2.0: [], 3.0: []}

    for i in range(n):
        curr_dt = pd.Timestamp(dates_np[i])
        if curr_dt in sup_by_date:
            for s in sup_by_date[curr_dt]:
                active.append({"price": float(s["price"]), "timeframe": s["timeframe"], "swept": False})
        if curr_dt < START:
            for sup in active:
                if not sup["swept"] and lows[i] < sup["price"]:
                    sup["swept"] = True
            continue
        for sup in list(active):
            if sup["swept"] or lows[i] >= sup["price"]:
                continue
            sup["swept"] = True
            support, tf = _resolve_effective_liquidity(float(sup["price"]), str(sup["timeframe"]), active)
            got = _try_setup(i, support, opens, highs, lows, closes, n)
            if got is None:
                continue
            eidx, entry, sl, _c1 = got
            for rr in (2.0, 3.0):
                exit_p, xidx, mfe, hit_sl, reason = _rr_exit(
                    eidx, entry, sl, rr, opens, highs, lows, n
                )
                raw[rr].append(
                    _pack(symbol, tf, nifty, support, eidx, xidx, entry, sl, exit_p, reason, 1, date_s, rr)
                )
                if hit_sl and mfe < 2.0:
                    sw = _next_sweep(xidx + 1, sl, lows, n)
                    if sw is None:
                        continue
                    got2 = _try_setup(sw, sl, opens, highs, lows, closes, n)
                    if got2 is None:
                        continue
                    e2, entry2, sl2, _ = got2
                    exit2, x2, _m2, _hs2, r2 = _rr_exit(e2, entry2, sl2, rr, opens, highs, lows, n)
                    raw[rr].append(
                        _pack(symbol, tf, nifty, sl, e2, x2, entry2, sl2, exit2, r2, 2, date_s, rr)
                    )

    out: dict[float, list[dict]] = {}
    for rr, lst in raw.items():
        lst.sort(key=lambda t: (t["Entry_Date"], -t["TF_Rank"]))
        kept, seen = [], set()
        for t in lst:
            k = t["Entry_Date"]
            if k in seen:
                continue
            seen.add(k)
            kept.append(t)
        out[rr] = kept
    return out


def _worker(symbol: str) -> dict[float, list[dict]]:
    try:
        return scan_ticker(symbol)
    except Exception:
        return {2.0: [], 3.0: []}


def _metrics(initial: float, final_eq: float, years: float) -> tuple[float, float]:
    if initial <= 0:
        return 0.0, 0.0
    ret = ((final_eq - initial) / initial) * 100.0
    if final_eq <= 0:
        return round(ret, 2), -100.0
    cagr = (((final_eq / initial) ** (1.0 / years)) - 1.0) * 100.0
    return round(ret, 2), round(cagr, 2)


def run_portfolio(df_trades: pd.DataFrame, capital: float, risk: float, rr: float, exp_name: str) -> dict:
    reports_dir = OUT_BASE / exp_name
    plots_dir = PLOT_BASE / exp_name
    reports_dir.mkdir(parents=True, exist_ok=True)
    plots_dir.mkdir(parents=True, exist_ok=True)
    rr_mode = f"1:{int(rr)}"

    df = df_trades.copy()
    df["Entry_Date"] = pd.to_datetime(df["Entry_Date"])
    by_entry: dict = defaultdict(list)
    for row in df.to_dict("records"):
        by_entry[row["Entry_Date"]].append(row)

    min_dt = pd.Timestamp("2010-01-01")
    max_dt = max(df["Entry_Date"].max(), pd.to_datetime(df["Exit_Date"]).max())
    cash = capital
    peak = capital
    max_dd = 0.0
    tx_id = 1
    trade_id = 1
    open_pos: dict[int, dict] = {}
    executed = 0
    total_tax_z = total_tax_f = total_gross = 0.0
    n_att1 = n_att2 = 0
    rows = [{
        "Transaction_ID": 1, "Trade_ID": 0, "Type": "DEPOSIT", "Date": "2010-01-01",
        "Ticker": "N/A", "Liquidity_Source": "N/A", "Support_Price": 0.0, "Quantity": 0,
        "Price": 0.0, "Total_Spend": 0.0, "Gross_PnL": 0.0, "Statutory_Taxes": 0.0,
        "Net_PnL": 0.0, "Return_Pct": 0.0, "Cash_Balance": cash,
        "Active_Position_Count": 0, "Holding_Equity_Value": 0.0,
        "Total_Portfolio_Value": cash, "Target_RR_Mode": rr_mode,
        "Outcome": "DEPOSIT", "Chart_PNG_URI": "N/A",
    }]
    tx_id = 2
    daily = []

    print(f"  Portfolio {exp_name}: Rs {capital:,.0f} / risk Rs {risk:,.0f} / {rr_mode}", flush=True)

    for day in pd.date_range(min_dt, max_dt, freq="D"):
        day_start = cash
        day_pnl = 0.0
        avail = cash
        cands = by_entry.get(day, [])
        cands.sort(key=lambda t: (-int(t.get("TF_Rank", 1)), -int(t.get("Nifty_Rank", 1))))
        for cand in cands:
            entry_p = float(cand["Entry_Price"])
            sl_p = float(cand["SL_Price"])
            rsk = entry_p - sl_p
            if rsk <= 0.05 or rsk > risk:
                continue
            qty = min(int(risk // rsk), int(avail // entry_p))
            if qty < 1:
                continue
            spend = round(entry_p * qty, 2)
            if spend > avail:
                continue
            cash = round(cash - spend, 2)
            avail = round(avail - spend, 2)
            executed += 1
            if int(cand.get("Attempt", 1)) == 2:
                n_att2 += 1
            else:
                n_att1 += 1
            open_pos[trade_id] = {
                "Ticker": cand["Ticker"],
                "Liquidity_Source": cand.get("Liquidity_Type", "Weekly"),
                "Support_Price": float(cand["Support_Price"]),
                "Entry_Price": entry_p,
                "Exit_Date": pd.Timestamp(cand["Exit_Date"]),
                "Exit_Price": float(cand["Exit_Price"]),
                "Quantity": qty,
                "Total_Spend": spend,
                "Outcome": cand.get("Outcome", "Failure"),
                "Attempt": int(cand.get("Attempt", 1)),
            }
            holding = sum(p["Total_Spend"] for p in open_pos.values())
            rows.append({
                "Transaction_ID": tx_id, "Trade_ID": trade_id, "Type": "BUY (ENTRY)",
                "Date": day.strftime("%Y-%m-%d"), "Ticker": cand["Ticker"],
                "Liquidity_Source": cand.get("Liquidity_Type", "Weekly"),
                "Support_Price": float(cand["Support_Price"]),
                "Quantity": qty, "Price": entry_p, "Total_Spend": spend,
                "Gross_PnL": 0.0, "Statutory_Taxes": 0.0, "Net_PnL": 0.0, "Return_Pct": 0.0,
                "Cash_Balance": cash, "Active_Position_Count": len(open_pos),
                "Holding_Equity_Value": holding, "Total_Portfolio_Value": cash + holding,
                "Target_RR_Mode": f"{rr_mode}|Att{int(cand.get('Attempt', 1))}",
                "Outcome": "OPEN", "Chart_PNG_URI": "N/A",
            })
            tx_id += 1
            trade_id += 1

        for tid, pos in list(open_pos.items()):
            if pos["Exit_Date"] > day:
                continue
            qty = pos["Quantity"]
            entry_p, exit_p = pos["Entry_Price"], pos["Exit_Price"]
            spend = pos["Total_Spend"]
            gross = round((exit_p - entry_p) * qty, 2)
            ch_z = calculate_indian_trade_charges(entry_p, exit_p, qty, 0.0)
            ch_f = calculate_indian_trade_charges(entry_p, exit_p, qty, 20.0)
            tax_z = round(ch_z["total_charges"], 2)
            tax_f = round(ch_f["total_charges"], 2)
            net = round(gross - tax_z, 2)
            total_gross += gross
            total_tax_z += tax_z
            total_tax_f += tax_f
            day_pnl += net
            cash = round(cash + spend + gross - tax_z, 2)
            del open_pos[tid]
            holding = sum(p["Total_Spend"] for p in open_pos.values())
            outcome = "Success" if net >= 0 else "Failure"
            rows.append({
                "Transaction_ID": tx_id, "Trade_ID": tid, "Type": "SELL (EXIT)",
                "Date": pos["Exit_Date"].strftime("%Y-%m-%d"), "Ticker": pos["Ticker"],
                "Liquidity_Source": pos["Liquidity_Source"], "Support_Price": pos["Support_Price"],
                "Quantity": qty, "Price": exit_p, "Total_Spend": spend,
                "Gross_PnL": gross, "Statutory_Taxes": tax_z, "Net_PnL": net,
                "Return_Pct": round((net / spend) * 100, 2) if spend else 0.0,
                "Cash_Balance": cash, "Active_Position_Count": len(open_pos),
                "Holding_Equity_Value": holding, "Total_Portfolio_Value": cash + holding,
                "Target_RR_Mode": f"{rr_mode}|Att{pos['Attempt']}",
                "Outcome": outcome, "Chart_PNG_URI": "N/A",
            })
            tx_id += 1

        port = cash + sum(p["Total_Spend"] for p in open_pos.values())
        peak = max(peak, port)
        max_dd = max(max_dd, ((peak - port) / peak * 100) if peak > 0 else 0)
        daily.append({
            "Date": day.strftime("%Y-%m-%d"),
            "Balance": cash,
            "Active_Positions": len(open_pos),
            "Daily_PnL": day_pnl,
            "Daily_Return_Pct": ((cash - day_start) / day_start * 100) if day_start else 0,
        })

    sells = [r for r in rows if r["Type"] == "SELL (EXIT)"]
    wins = sum(1 for r in sells if r["Net_PnL"] >= 0)
    win_rate = (wins / len(sells) * 100) if sells else 0.0
    years = max((max_dt - min_dt).days / 365.25, 0.01)
    gross_eq = round(capital + total_gross, 2)
    net_z = cash
    net_f = round(net_z - (total_tax_f - total_tax_z), 2)
    g_ret, g_cagr = _metrics(capital, gross_eq, years)
    z_ret, z_cagr = _metrics(capital, net_z, years)
    f_ret, f_cagr = _metrics(capital, net_f, years)

    df_stmt = pd.DataFrame(rows)
    df_daily = pd.DataFrame(daily)
    df_stmt.to_csv(reports_dir / "Swing_Strategy_Account_Statement.csv", index=False)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Account Statement"
    headers = list(df_stmt.columns)
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")
    for _, r in df_stmt.iterrows():
        ws.append(list(r))
    ws_sum = wb.create_sheet("Performance Summary")
    ws_sum.append(["Performance Metric", "Gross (BEFORE TAX)", "Net (Zerodha)", "Net (Flat Rs 20)"])
    period = f"{min_dt.strftime('%Y-%m-%d')} to {max_dt.strftime('%Y-%m-%d')}"
    for row in [
        ("Strategy Rule", RULE, RULE, RULE),
        ("Target RR Mode", rr_mode, rr_mode, rr_mode),
        ("Initial Capital", f"Rs {capital:,.2f}", f"Rs {capital:,.2f}", f"Rs {capital:,.2f}"),
        ("Fixed Risk Cap / Trade", f"Rs {risk:,.2f}", f"Rs {risk:,.2f}", f"Rs {risk:,.2f}"),
        ("Final Portfolio Equity", f"Rs {gross_eq:,.2f}", f"Rs {net_z:,.2f}", f"Rs {net_f:,.2f}"),
        ("Total Net Profit (INR)", f"Rs {gross_eq - capital:,.2f}", f"Rs {net_z - capital:,.2f}", f"Rs {net_f - capital:,.2f}"),
        ("Total Return (%)", f"{g_ret:+.2f}%", f"{z_ret:+.2f}%", f"{f_ret:+.2f}%"),
        ("CAGR (%)", f"{g_cagr:+.2f}%", f"{z_cagr:+.2f}%", f"{f_cagr:+.2f}%"),
        ("Executed Trades Count", f"{executed:,}", f"{executed:,}", f"{executed:,}"),
        ("Attempt1 / Attempt2", f"{n_att1} / {n_att2}", f"{n_att1} / {n_att2}", f"{n_att1} / {n_att2}"),
        ("Win Rate (%)", f"{win_rate:.2f}%", f"{win_rate:.2f}%", f"{win_rate:.2f}%"),
        ("Max Drawdown (%)", f"{max_dd:.2f}%", f"{max_dd:.2f}%", f"{max_dd:.2f}%"),
        ("Total Statutory Taxes Paid", "Rs 0.00", f"Rs {total_tax_z:,.2f}", f"Rs {total_tax_f:,.2f}"),
        ("Backtest Period", period, period, period),
    ]:
        ws_sum.append(list(row))
    wb.save(reports_dir / "Swing_Strategy_Account_Statement.xlsx")
    generate_all_visualizations(df_daily, plots_dir, reports_dir, exp_title=exp_name)
    print(
        f"  DONE {exp_name} | Final Rs {net_z:,.2f} | Net CAGR {z_cagr:+.2f}% | "
        f"Trades {executed:,} (att2 {n_att2:,}) | Win {win_rate:.1f}% | DD {max_dd:.1f}%",
        flush=True,
    )
    return {
        "Experiment": exp_name,
        "RR": rr_mode,
        "Capital": capital,
        "Risk": risk,
        "Final_Zerodha": net_z,
        "Gross_Equity": gross_eq,
        "Net_Return_Pct": z_ret,
        "CAGR_Pct": z_cagr,
        "Gross_CAGR_Pct": g_cagr,
        "Fyers_CAGR_Pct": f_cagr,
        "Executed": executed,
        "Attempt1": n_att1,
        "Attempt2": n_att2,
        "Win_Rate_Pct": round(win_rate, 2),
        "Max_DD_Pct": round(max_dd, 2),
    }


def main():
    tickers = sorted({p.name.split("_1d.csv")[0] for p in DATA_DAILY_DIR.glob("*_1d.csv")})
    print(f"{RULE}\nScanning {len(tickers):,} tickers for RR 1:2 and 1:3 ...", flush=True)
    by_rr: dict[float, list[dict]] = {2.0: [], 3.0: []}
    done = 0
    with ProcessPoolExecutor(max_workers=8) as ex:
        futs = {ex.submit(_worker, t): t for t in tickers}
        for fut in as_completed(futs):
            got = fut.result()
            by_rr[2.0].extend(got.get(2.0, []))
            by_rr[3.0].extend(got.get(3.0, []))
            done += 1
            if done % 200 == 0 or done == len(tickers):
                print(
                    f"  • {done:,}/{len(tickers):,} | 1:2 setups {len(by_rr[2.0]):,} | 1:3 setups {len(by_rr[3.0]):,}",
                    flush=True,
                )

    OUT_BASE.mkdir(parents=True, exist_ok=True)
    summaries = []
    for sc in SCENARIOS:
        df = pd.DataFrame(by_rr[sc["rr"]])
        if df.empty:
            print(f"No trades for RR {sc['rr']}", flush=True)
            continue
        summaries.append(run_portfolio(df, sc["capital"], sc["risk"], sc["rr"], sc["name"]))

    cmp = pd.DataFrame(summaries)
    cmp_path = OUT_BASE / "Retry2_1to2_vs_1to3_Comparison.csv"
    cmp.to_csv(cmp_path, index=False)
    print("\n" + cmp.to_string(index=False), flush=True)
    print(f"\nSaved {cmp_path}", flush=True)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
