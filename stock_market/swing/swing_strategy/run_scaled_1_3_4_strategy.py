"""
Scaled exits 1:1 / 1:3 / 1:4 with TF-then-Nifty selection.

Rules:
- On sweep, if next support below is within 5%, use the lower level
- C1: green candle, close below liquidity
- C2: High > C1 high before C1 low breaks; entry = C1_High x 1.001 (TouchPlannedEntry)
- SL = C1 low
- Exits: 33% @ 1:1, 33% @ 1:3, 34% @ 1:4
- Selection: Yearly > Monthly > Weekly, then Nifty 50 > 100 > 250 > Other
- Default capital Rs 50,000 / Rs 1,000 risk per trade
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

EXP_NAME = "Exp_50k_1k_1to1_1to3_1to4"
REPORTS_BASE = BASE_DIR / "Reports" / "Scaled_1_3_4_TFNifty"
PLOTS_BASE = BASE_DIR / "Plots" / "Scaled_1_3_4_TFNifty"
CACHE_CSV = REPORTS_BASE / "Scaled_1_3_4_Trades.csv"
RR_MODE = "33% 1:1 / 33% 1:3 / 34% 1:4"

from src.analysis.indian_brokerage_calculator import calculate_indian_trade_charges
from swing_strategy.tiered_liquidity_strategy_engine import build_trades_dataset
from swing_strategy.visualizer import generate_all_visualizations

TF_RANK = {"Yearly": 3, "Monthly": 2, "Weekly": 1}
NIFTY_RANK = {"Nifty 50": 4, "Nifty 100": 3, "Nifty 250": 2, "Other": 1}

LEG_RR_LABEL = {
    "TP1": "1:1",
    "TP1_GAP": "1:1",
    "TP2": "1:3",
    "TP2_GAP": "1:3",
    "TP3": "1:4",
    "TP3_GAP": "1:4",
    "SL": "SL",
    "SL_GAP": "SL",
    "EOD": "EOD",
}


def _selection_sort_key(cand: dict) -> tuple:
    """Yearly > Monthly > Weekly, then Nifty 50 > 100 > 250 > Other."""
    return (
        -TF_RANK.get(cand.get("Liquidity_Type", "Weekly"), 0),
        -NIFTY_RANK.get(cand.get("Index_Membership", "Other"), 0),
    )


def _metrics(initial: float, final_eq: float, num_years: float) -> tuple[float, float]:
    if initial <= 0:
        return 0.0, 0.0
    ret = ((final_eq - initial) / initial) * 100.0
    if final_eq <= 0:
        return round(ret, 2), -100.0
    cagr = (((final_eq / initial) ** (1.0 / num_years)) - 1.0) * 100.0
    return round(ret, 2), round(cagr, 2)


def run_portfolio_backtest(
    df_trades: pd.DataFrame,
    initial_capital: float = 50_000.0,
    risk_per_trade: float = 1_000.0,
) -> dict:
    reports_dir = REPORTS_BASE / EXP_NAME
    plots_dir = PLOTS_BASE / EXP_NAME
    reports_dir.mkdir(parents=True, exist_ok=True)
    plots_dir.mkdir(parents=True, exist_ok=True)

    if df_trades.empty:
        return {"error": "no trades"}

    df = df_trades.copy()
    df["Entry_Date"] = pd.to_datetime(df["Entry_Date"])

    trades_by_entry: dict = {}
    for row in df.to_dict("records"):
        trades_by_entry.setdefault(row["Entry_Date"], []).append(row)

    min_dt = pd.to_datetime("2010-01-01")
    max_dt = max(df["Entry_Date"].max(), pd.to_datetime(df["Exit_Date"]).max())
    all_days = pd.date_range(min_dt, max_dt, freq="D")

    current_cash = initial_capital
    peak_equity = initial_capital
    max_dd = 0.0
    tx_id = 1
    trade_id = 1

    statement_rows = []
    daily_equity = []
    daily_entry_report = []
    open_positions: dict[int, dict] = {}

    executed = 0
    total_taxes_z = 0.0
    total_taxes_f = 0.0
    total_gross = 0.0

    statement_rows.append({
        "Transaction_ID": tx_id, "Trade_ID": 0, "Type": "DEPOSIT", "Date": "2010-01-01",
        "Ticker": "N/A", "Liquidity_Source": "N/A", "Support_Price": 0.0, "Quantity": 0,
        "Price": 0.0, "Total_Spend": 0.0, "Gross_PnL": 0.0, "Statutory_Taxes": 0.0,
        "Net_PnL": 0.0, "Return_Pct": 0.0, "Cash_Balance": current_cash,
        "Active_Position_Count": 0, "Holding_Equity_Value": 0.0,
        "Total_Portfolio_Value": current_cash, "Target_RR_Mode": RR_MODE,
        "Outcome": "DEPOSIT", "Chart_PNG_URI": "N/A",
    })
    tx_id += 1

    print(
        f"Portfolio: Rs {initial_capital:,.0f} | risk Rs {risk_per_trade:,.0f} | {RR_MODE}",
        flush=True,
    )
    print("Selection: Yearly > Monthly > Weekly, then Nifty 50 > 100 > 250 > Other", flush=True)

    for curr_dt in all_days:
        day_start = current_cash
        day_pnl = 0.0

        signals_today = trades_by_entry.get(curr_dt, [])
        signals_today.sort(key=_selection_sort_key)

        attempted = len(signals_today)
        entered = 0
        skipped_cash = 0
        entered_tickers = []

        available = current_cash
        for cand in signals_today:
            entry_p = float(cand["Entry_Price"])
            sl_p = float(cand["SL_Price"])
            risk = max(0.05, entry_p - sl_p)
            qty = int(risk_per_trade / risk)
            if qty < 1:
                skipped_cash += 1
                continue
            pos_val = round(entry_p * qty, 2)
            if pos_val > available:
                skipped_cash += 1
                continue

            current_cash = round(current_cash - pos_val, 2)
            available = round(available - pos_val, 2)
            executed += 1
            entered += 1
            entered_tickers.append(cand["Ticker"])

            legs = json.loads(cand["Exit_Legs_JSON"])
            open_positions[trade_id] = {
                "Trade_ID": trade_id,
                "Ticker": cand["Ticker"],
                "Liquidity_Source": cand.get("Liquidity_Type", "Weekly"),
                "Support_Price": float(cand["Support_Price"]),
                "C1_Date": cand["C1_Date"],
                "Entry_Date": curr_dt.strftime("%Y-%m-%d"),
                "Entry_Price": entry_p,
                "SL_Price": sl_p,
                "Quantity": qty,
                "Total_Spend": pos_val,
                "Remaining_Qty": qty,
                "Remaining_Spend": pos_val,
                "Legs": legs,
                "Legs_Done": 0,
                "First_Exit": True,
            }

            holding = sum(p["Remaining_Spend"] for p in open_positions.values())
            statement_rows.append({
                "Transaction_ID": tx_id, "Trade_ID": trade_id, "Type": "BUY (ENTRY)",
                "Date": curr_dt.strftime("%Y-%m-%d"), "Ticker": cand["Ticker"],
                "Liquidity_Source": cand["Liquidity_Type"],
                "Support_Price": float(cand["Support_Price"]),
                "Quantity": qty, "Price": entry_p, "Total_Spend": pos_val,
                "Gross_PnL": 0.0, "Statutory_Taxes": 0.0, "Net_PnL": 0.0, "Return_Pct": 0.0,
                "Cash_Balance": current_cash, "Active_Position_Count": len(open_positions),
                "Holding_Equity_Value": holding, "Total_Portfolio_Value": current_cash + holding,
                "Target_RR_Mode": RR_MODE, "Outcome": "OPEN", "Chart_PNG_URI": "N/A",
            })
            tx_id += 1
            trade_id += 1

        daily_entry_report.append({
            "Date": curr_dt.strftime("%Y-%m-%d"),
            "Signals_Matching_Criteria": attempted,
            "Entries_Attempted": attempted,
            "Entries_Executed": entered,
            "Skipped_Insufficient_Cash": skipped_cash,
            "Tickers_Entered": ", ".join(entered_tickers) if entered_tickers else "",
        })

        for tid, pos in list(open_positions.items()):
            while pos["Legs_Done"] < len(pos["Legs"]):
                leg = pos["Legs"][pos["Legs_Done"]]
                leg_dt = pd.to_datetime(leg["date"])
                if leg_dt > curr_dt:
                    break

                sell_qty = max(1, int(round(pos["Quantity"] * leg["qty_pct"])))
                sell_qty = min(sell_qty, pos["Remaining_Qty"])
                if sell_qty <= 0:
                    pos["Legs_Done"] += 1
                    continue

                entry_p = pos["Entry_Price"]
                exit_p = float(leg["price"])
                spend_frac = sell_qty / pos["Quantity"]
                spend_part = round(pos["Total_Spend"] * spend_frac, 2)

                gross = round((exit_p - entry_p) * sell_qty, 2)
                total_gross += gross
                ch_z = calculate_indian_trade_charges(entry_p, exit_p, sell_qty, flat_brokerage_per_order=0.0)
                ch_f = calculate_indian_trade_charges(entry_p, exit_p, sell_qty, flat_brokerage_per_order=20.0)
                tax_z = round(ch_z["total_charges"], 2)
                tax_f = round(ch_f["total_charges"], 2)
                total_taxes_z += tax_z
                total_taxes_f += tax_f
                net = round(gross - tax_z, 2)
                day_pnl += net
                ret_pct = round((net / spend_part) * 100.0, 2) if spend_part > 0 else 0.0

                current_cash = round(current_cash + spend_part + gross - tax_z, 2)
                pos["Remaining_Qty"] -= sell_qty
                pos["Remaining_Spend"] = round(pos["Remaining_Spend"] - spend_part, 2)
                pos["Legs_Done"] += 1
                pos["First_Exit"] = False

                if pos["Remaining_Qty"] <= 0:
                    del open_positions[tid]

                outcome = "Success" if net >= 0 else "Failure"
                leg_name = str(leg.get("leg", "EXIT"))
                holding = sum(p["Remaining_Spend"] for p in open_positions.values())

                statement_rows.append({
                    "Transaction_ID": tx_id, "Trade_ID": tid, "Type": "SELL (EXIT)",
                    "Date": leg_dt.strftime("%Y-%m-%d"), "Ticker": pos["Ticker"],
                    "Liquidity_Source": pos["Liquidity_Source"], "Support_Price": pos["Support_Price"],
                    "Quantity": sell_qty, "Price": exit_p, "Total_Spend": spend_part,
                    "Gross_PnL": gross, "Statutory_Taxes": tax_z, "Net_PnL": net,
                    "Return_Pct": ret_pct, "Cash_Balance": current_cash,
                    "Active_Position_Count": len(open_positions),
                    "Holding_Equity_Value": holding, "Total_Portfolio_Value": current_cash + holding,
                    "Target_RR_Mode": LEG_RR_LABEL.get(leg_name, RR_MODE),
                    "Outcome": outcome, "Chart_PNG_URI": "N/A",
                })
                tx_id += 1

                if tid not in open_positions:
                    break

        port = current_cash + sum(p["Remaining_Spend"] for p in open_positions.values())
        if port > peak_equity:
            peak_equity = port
        max_dd = max(max_dd, ((peak_equity - port) / peak_equity * 100) if peak_equity > 0 else 0)

        daily_equity.append({
            "Date": curr_dt.strftime("%Y-%m-%d"),
            "Balance": current_cash,
            "Active_Positions": len(open_positions),
            "Daily_PnL": day_pnl,
            "Daily_Return_Pct": ((current_cash - day_start) / day_start * 100) if day_start > 0 else 0,
        })

    sell_rows = [r for r in statement_rows if r["Type"] == "SELL (EXIT)"]
    trade_pnl: dict = {}
    for r in sell_rows:
        trade_pnl[r["Trade_ID"]] = trade_pnl.get(r["Trade_ID"], 0.0) + r["Net_PnL"]
    wins = sum(1 for v in trade_pnl.values() if v >= 0)
    win_rate = (wins / len(trade_pnl) * 100) if trade_pnl else 0.0

    df_stmt = pd.DataFrame(statement_rows)
    df_daily = pd.DataFrame(daily_equity)
    df_entry = pd.DataFrame(daily_entry_report)

    num_years = max((max_dt - min_dt).days / 365.25, 0.01)
    net_eq_z = current_cash
    gross_eq = round(initial_capital + total_gross, 2)
    net_eq_f = round(net_eq_z - (total_taxes_f - total_taxes_z), 2)

    gross_ret, gross_cagr = _metrics(initial_capital, gross_eq, num_years)
    net_ret, net_cagr = _metrics(initial_capital, net_eq_z, num_years)
    fyers_ret, fyers_cagr = _metrics(initial_capital, net_eq_f, num_years)

    csv_path = reports_dir / "Swing_Strategy_Account_Statement.csv"
    xlsx_path = reports_dir / "Swing_Strategy_Account_Statement.xlsx"
    entry_path = reports_dir / "Daily_Entry_Report.csv"
    df_stmt.to_csv(csv_path, index=False)
    df_entry.to_csv(entry_path, index=False)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    buy_fill = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    sell_win_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    sell_loss_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Account Statement"
    headers = list(df_stmt.columns)
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, r in df_stmt.iterrows():
        row_num = r_idx + 2
        tx_type = r["Type"]
        pnl = r["Net_PnL"]
        if str(tx_type).startswith("BUY"):
            fill = buy_fill
        elif str(tx_type).startswith("SELL") and pnl >= 0:
            fill = sell_win_fill
        elif str(tx_type).startswith("SELL"):
            fill = sell_loss_fill
        else:
            fill = None
        for c_idx, val in enumerate(r, start=1):
            cell = ws.cell(row=row_num, column=c_idx, value=val)
            if fill is not None:
                cell.fill = fill
            cell.alignment = Alignment(horizontal="center" if c_idx in (1, 2, 3, 4, 5, 6, 19, 20) else "right")

    ws_sum = wb.create_sheet("Performance Summary")
    ws_sum.append(["Performance Metric", "Gross (BEFORE TAX)", "Net (Zerodha)", "Net (Flat Rs 20)"])
    for col in range(1, 5):
        c = ws_sum.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    period = f"{min_dt.strftime('%Y-%m-%d')} to {max_dt.strftime('%Y-%m-%d')}"
    for row in [
        ("Strategy Rule", "5% lower-liq + C1 green close-below + TouchPlannedEntry",
         "5% lower-liq + C1 green close-below + TouchPlannedEntry",
         "5% lower-liq + C1 green close-below + TouchPlannedEntry"),
        ("Selection", "Yearly>Monthly>Weekly then Nifty 50>100>250>Other",
         "Yearly>Monthly>Weekly then Nifty 50>100>250>Other",
         "Yearly>Monthly>Weekly then Nifty 50>100>250>Other"),
        ("Target RR Mode", RR_MODE, RR_MODE, RR_MODE),
        ("Initial Capital", f"Rs {initial_capital:,.2f}", f"Rs {initial_capital:,.2f}", f"Rs {initial_capital:,.2f}"),
        ("Fixed Risk Cap / Trade", f"Rs {risk_per_trade:,.2f}", f"Rs {risk_per_trade:,.2f}", f"Rs {risk_per_trade:,.2f}"),
        ("Final Portfolio Equity", f"Rs {gross_eq:,.2f}", f"Rs {net_eq_z:,.2f}", f"Rs {net_eq_f:,.2f}"),
        ("Total Net Profit (INR)", f"Rs {gross_eq - initial_capital:,.2f}",
         f"Rs {net_eq_z - initial_capital:,.2f}", f"Rs {net_eq_f - initial_capital:,.2f}"),
        ("Total Return (%)", f"{gross_ret:+.2f}%", f"{net_ret:+.2f}%", f"{fyers_ret:+.2f}%"),
        ("CAGR (%)", f"{gross_cagr:+.2f}%", f"{net_cagr:+.2f}%", f"{fyers_cagr:+.2f}%"),
        ("Executed Trades Count", f"{executed:,}", f"{executed:,}", f"{executed:,}"),
        ("Win Rate (%)", f"{win_rate:.2f}%", f"{win_rate:.2f}%", f"{win_rate:.2f}%"),
        ("Max Drawdown (%)", f"{max_dd:.2f}%", f"{max_dd:.2f}%", f"{max_dd:.2f}%"),
        ("Total Statutory Taxes Paid", "Rs 0.00", f"Rs {total_taxes_z:,.2f}", f"Rs {total_taxes_f:,.2f}"),
        ("Backtest Period", period, period, period),
    ]:
        ws_sum.append(list(row))
    ws_sum.column_dimensions["A"].width = 38
    ws_sum.column_dimensions["B"].width = 56
    ws_sum.column_dimensions["C"].width = 56
    ws_sum.column_dimensions["D"].width = 56
    wb.save(xlsx_path)

    generate_all_visualizations(df_daily, plots_dir, reports_dir, exp_title=EXP_NAME)

    print(
        f"DONE | Final Rs {net_eq_z:,.2f} | Net CAGR {net_cagr:+.2f}% | "
        f"Gross CAGR {gross_cagr:+.2f}% | Trades {executed:,} | Win {win_rate:.2f}% | DD {max_dd:.2f}%",
        flush=True,
    )
    print(f"Statement: {xlsx_path}", flush=True)
    print(f"Daily entry report: {entry_path}", flush=True)

    return {
        "Experiment": EXP_NAME,
        "Initial_Capital": initial_capital,
        "Risk_Cap": risk_per_trade,
        "Final_Equity": net_eq_z,
        "Gross_Final_Equity": gross_eq,
        "Net_Return_Pct": net_ret,
        "CAGR_Pct": net_cagr,
        "Gross_CAGR_Pct": gross_cagr,
        "Fyers_CAGR_Pct": fyers_cagr,
        "Executed_Trades": executed,
        "Win_Rate_Pct": round(win_rate, 2),
        "Max_Drawdown_Pct": round(max_dd, 2),
        "Taxes_Zerodha": round(total_taxes_z, 2),
    }


def main():
    parser = argparse.ArgumentParser(description="Scaled 1:1/1:3/1:4 TF-then-Nifty strategy")
    parser.add_argument("--capital", type=float, default=50_000.0)
    parser.add_argument("--risk", type=float, default=1_000.0)
    parser.add_argument("--rescan", action="store_true")
    args = parser.parse_args()

    print("=" * 72, flush=True)
    print("  SCALED EXITS 1:1 / 1:3 / 1:4  |  TF then Nifty selection", flush=True)
    print(f"  Capital Rs {args.capital:,.0f}  |  Risk Rs {args.risk:,.0f}", flush=True)
    print("=" * 72, flush=True)

    df = build_trades_dataset(
        force_rescan=args.rescan,
        rr1=1.0,
        rr2=3.0,
        rr3=4.0,
        cache_csv=CACHE_CSV,
    )
    if df.empty:
        print("No setups found. Aborting.", flush=True)
        return
    summary = run_portfolio_backtest(df, initial_capital=args.capital, risk_per_trade=args.risk)
    print(summary, flush=True)


if __name__ == "__main__":
    main()
