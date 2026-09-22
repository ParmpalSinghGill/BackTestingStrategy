"""
Pure No-ML Strategy Benchmark Suite (Fixed 1:2 RR & Fixed 1:3 RR)

Evaluates quantitative swing strategy performance WITHOUT any Machine Learning filtering.
Takes ALL valid trade setups (Scenario 1: Green & Close > C1 High) across 2010 to 2026.

Scenarios Tested (Initial Capital Rs 100,000):
1. Pure No-ML Fixed 1:2 RR (Rs 1,000 Risk Cap)
2. Pure No-ML Fixed 1:2 RR (Rs 500 Risk Cap)
3. Pure No-ML Fixed 1:3 RR (Rs 1,000 Risk Cap)
4. Pure No-ML Fixed 1:3 RR (Rs 500 Risk Cap)

Follows Guide/Realistic_guide.md and Guide/Account_Statement_guide.md.
"""

import os
import sys
from pathlib import Path
import pandas as pd
import numpy as np

# Force UTF-8 encoding for stdout on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

REPORTS_DIR = BASE_DIR / "Reports"
NO_ML_DIR = REPORTS_DIR / "No_ML_Benchmark_Suite"
PLOTS_DIR = BASE_DIR / "Plots" / "No_ML_Benchmark_Suite"

os.makedirs(NO_ML_DIR, exist_ok=True)
os.makedirs(PLOTS_DIR, exist_ok=True)

from swing_strategy.generate_statement import generate_swing_strategy_statement

dataset_path = REPORTS_DIR / "Exact_True_6Class_Trade_Features_Dataset.csv"
if not dataset_path.exists():
    dataset_path = REPORTS_DIR / "Four_Class_Trade_Features_Dataset.csv"

df = pd.read_csv(dataset_path)

df["C2_Date"] = pd.to_datetime(df["C2_Date"])
df = df.sort_values("C2_Date").reset_index(drop=True)

sc_col = "Scenario" if "Scenario" in df.columns else "Scenario_1to2"
df_sc1 = df[df[sc_col] == "Scenario 1 (Green & Close > C1 High)"].copy()
df_sc1 = df_sc1.sort_values("C2_Date").reset_index(drop=True)

print(f"=== Pure No-ML Strategy Suite (Total Setups Available: {len(df_sc1):,}) ===", flush=True)

# 1. No-ML Fixed 1:2 RR (Risk Cap Rs 1,000)
df_1to2 = df_sc1.copy()
df_1to2["ML_RR_Choice"] = "1:2"
df_1to2["ML_Prediction"] = "Enter"

exp1 = "No_ML_Fixed_1to2_Risk1k"
sum1 = generate_swing_strategy_statement(
    initial_deposit=100000.0, max_charts_to_generate=0, max_risk_per_trade=1000.0,
    df_acc=df_1to2, custom_reports_dir=NO_ML_DIR / exp1, custom_plots_dir=PLOTS_DIR / exp1, exp_name=exp1
)

# 2. No-ML Fixed 1:2 RR (Risk Cap Rs 500)
exp2 = "No_ML_Fixed_1to2_Risk500"
sum2 = generate_swing_strategy_statement(
    initial_deposit=100000.0, max_charts_to_generate=0, max_risk_per_trade=500.0,
    df_acc=df_1to2, custom_reports_dir=NO_ML_DIR / exp2, custom_plots_dir=PLOTS_DIR / exp2, exp_name=exp2
)

# 3. No-ML Fixed 1:3 RR (Risk Cap Rs 1,000)
df_1to3 = df_sc1.copy()
df_1to3["ML_RR_Choice"] = "1:3"
df_1to3["ML_Prediction"] = "Enter"

exp3 = "No_ML_Fixed_1to3_Risk1k"
sum3 = generate_swing_strategy_statement(
    initial_deposit=100000.0, max_charts_to_generate=0, max_risk_per_trade=1000.0,
    df_acc=df_1to3, custom_reports_dir=NO_ML_DIR / exp3, custom_plots_dir=PLOTS_DIR / exp3, exp_name=exp3
)

# 4. No-ML Fixed 1:3 RR (Risk Cap Rs 500)
exp4 = "No_ML_Fixed_1to3_Risk500"
sum4 = generate_swing_strategy_statement(
    initial_deposit=100000.0, max_charts_to_generate=0, max_risk_per_trade=500.0,
    df_acc=df_1to3, custom_reports_dir=NO_ML_DIR / exp4, custom_plots_dir=PLOTS_DIR / exp4, exp_name=exp4
)

no_ml_summary = [
    {
        "Strategy Variant": "Pure No-ML Fixed 1:2 RR (Rs 1,000 Risk)",
        "Target RR": "1:2",
        "Risk Cap": "Rs 1,000",
        "Final Equity": f"Rs {sum1['Final_Equity']:,.2f}",
        "Net Return (%)": f"+{sum1['Total_Net_Return_Pct']:,.2f}%",
        "CAGR (%)": f"{sum1['CAGR_Pct']:.2f}%",
        "Executed Trades": f"{sum1['Executed_Trades']:,}",
        "Win Rate (%)": f"{sum1['Win_Rate_Pct']:.2f}%",
        "Max Drawdown (%)": f"{sum1['Max_Drawdown_Pct']:.2f}%",
        "Taxes Paid": f"Rs {sum1['Total_Taxes_Paid']:,.2f}"
    },
    {
        "Strategy Variant": "Pure No-ML Fixed 1:2 RR (Rs 500 Risk)",
        "Target RR": "1:2",
        "Risk Cap": "Rs 500",
        "Final Equity": f"Rs {sum2['Final_Equity']:,.2f}",
        "Net Return (%)": f"+{sum2['Total_Net_Return_Pct']:,.2f}%",
        "CAGR (%)": f"{sum2['CAGR_Pct']:.2f}%",
        "Executed Trades": f"{sum2['Executed_Trades']:,}",
        "Win Rate (%)": f"{sum2['Win_Rate_Pct']:.2f}%",
        "Max Drawdown (%)": f"{sum2['Max_Drawdown_Pct']:.2f}%",
        "Taxes Paid": f"Rs {sum2['Total_Taxes_Paid']:,.2f}"
    },
    {
        "Strategy Variant": "Pure No-ML Fixed 1:3 RR (Rs 1,000 Risk)",
        "Target RR": "1:3",
        "Risk Cap": "Rs 1,000",
        "Final Equity": f"Rs {sum3['Final_Equity']:,.2f}",
        "Net Return (%)": f"+{sum3['Total_Net_Return_Pct']:,.2f}%",
        "CAGR (%)": f"{sum3['CAGR_Pct']:.2f}%",
        "Executed Trades": f"{sum3['Executed_Trades']:,}",
        "Win Rate (%)": f"{sum3['Win_Rate_Pct']:.2f}%",
        "Max Drawdown (%)": f"{sum3['Max_Drawdown_Pct']:.2f}%",
        "Taxes Paid": f"Rs {sum3['Total_Taxes_Paid']:,.2f}"
    },
    {
        "Strategy Variant": "Pure No-ML Fixed 1:3 RR (Rs 500 Risk)",
        "Target RR": "1:3",
        "Risk Cap": "Rs 500",
        "Final Equity": f"Rs {sum4['Final_Equity']:,.2f}",
        "Net Return (%)": f"+{sum4['Total_Net_Return_Pct']:,.2f}%",
        "CAGR (%)": f"{sum4['CAGR_Pct']:.2f}%",
        "Executed Trades": f"{sum4['Executed_Trades']:,}",
        "Win Rate (%)": f"{sum4['Win_Rate_Pct']:.2f}%",
        "Max Drawdown (%)": f"{sum4['Max_Drawdown_Pct']:.2f}%",
        "Taxes Paid": f"Rs {sum4['Total_Taxes_Paid']:,.2f}"
    }
]

print("\n==========================================================================")
print("          PURE NO-ML STRATEGY BENCHMARK RESULTS (100k Initial Deposit)")
print("==========================================================================")
df_no_ml_sum = pd.DataFrame(no_ml_summary)
print(df_no_ml_sum.to_string(index=False))
print("==========================================================================\n")

csv_path = NO_ML_DIR / "Pure_No_ML_Strategy_Comparison.csv"
excel_path = NO_ML_DIR / "Pure_No_ML_Strategy_Comparison.xlsx"

df_no_ml_sum.to_csv(csv_path, index=False)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "No-ML Strategy Comparison"

headers = list(df_no_ml_sum.columns)
ws.append(headers)

header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for _, r in df_no_ml_sum.iterrows():
    ws.append(list(r))

wb.save(excel_path)
print(f"Saved Pure No-ML Strategy Comparison Excel & CSV!")
