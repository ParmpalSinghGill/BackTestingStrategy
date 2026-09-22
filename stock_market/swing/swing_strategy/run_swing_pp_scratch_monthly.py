"""Monthly-start vintages of Swing_PP C3-close scratch (isolated 1:2).

Fresh Rs 50,000 on the 1st of each month from 2005-01-01 through last data,
held to the last exit. Same live book: C3 open, sweep SL, Meta_P>=0.48 top 16,
2% equity, last-50 paper fail% >= 74%. Extra exit: on the entry bar (C3) only,
if SL and 1:2 are not tagged and C3 closes below C1 high, sell at C3 close.

21-col statements use MTM holding (qty x close) per Guide/Account_Statement_guide.md.
Does not write Swing_low / Swing_Live. Not a frozen Rs 500 book.
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from swing_strategy.run_2pct_50_25_hunt import CAPITAL
from swing_strategy.run_c2_cross_rr_hunt import exit_path, locate_c1
from swing_strategy.run_pred_paper_gate import exits_by_day, load_pred_list
from swing_strategy.run_swing_pp_rr_statement import FEAT_N, simulate_rr_book
from swing_strategy.run_swing_pp_statement import CSV_COLS, _write_excel
from swing_strategy.statement_mtm import CloseCache, index_trade_charts
from swing_strategy.tiered_liquidity_strategy_engine import _load_daily
from swing_strategy.visualizer import generate_all_visualizations

OUT = BASE / "Reports" / "SwingPP_Scratch_Monthly"
PLOTS = BASE / "Plots" / "SwingPP_Scratch_Monthly"
CHARTS = PLOTS / "trade_charts"
CACHE = BASE / "Reports" / "SwingLow_OldLiquidity" / "Pred_fairBE_scratch_c3.parquet"
MASTER = OUT / "Master_Experiments_Comparison.xlsx"
CSV = OUT / "Monthly_Start_CAGR.csv"
LOG = OUT / "Monthly_Start_CAGR.json"
HEATMAP = PLOTS / "Vintage_CAGR_Heatmap.png"
RULE = (
    "Swing_PP scratch | C3 open, sweep SL x0.99, 1:2 fair BE | "
    "entry-bar scratch if C3 close < C1 high | Meta_P>=0.48 top16 | "
    "2% equity | paper-gate last 50 fail%>=74%"
)
CHART_REL_PREFIX = "../../../Plots/SwingPP_Scratch_Monthly/trade_charts"


def rewrite_scratch(base: pd.DataFrame) -> pd.DataFrame:
    """Keep original Entry_Price / SL_Price; rewrite exit with C3-close scratch."""
    if CACHE.exists():
        df = pd.read_parquet(CACHE)
        df["Entry_Date"] = pd.to_datetime(df["Entry_Date"])
        df["Exit_Date"] = pd.to_datetime(df["Exit_Date"])
        print(f"[scratch] cache {len(df):,}  {CACHE.name}", flush=True)
        return df

    cache: dict = {}
    rows: list[dict] = []
    n_scratch = 0
    miss_c1 = 0
    recs = base.to_dict("records")
    print(f"[scratch] rewrite {len(recs):,} names (keep entry/SL)", flush=True)
    for i, rec in enumerate(recs):
        if i and i % 1500 == 0:
            print(f"  {i:,}/{len(recs):,}  scratch={n_scratch:,}  miss_c1={miss_c1:,}", flush=True)
        out = dict(rec)
        out["Target_RR_Mode"] = "1:2"
        out["chosen_rr"] = 2
        sym = rec["Ticker"]
        if sym not in cache:
            df = _load_daily(sym)
            if df is None or "Date" not in getattr(df, "columns", []):
                cache[sym] = None
            else:
                d = df.copy()
                d["Date"] = pd.to_datetime(d["Date"])
                cache[sym] = d.sort_values("Date").reset_index(drop=True)
        df = cache[sym]
        if df is None:
            rows.append(out)
            continue
        edt = pd.Timestamp(rec["Entry_Date"]).normalize()
        dates = pd.to_datetime(df["Date"]).dt.normalize()
        idx = int(dates.searchsorted(edt))
        if idx >= len(df) or dates.iloc[idx] != edt:
            rows.append(out)
            continue
        opens = df["Open"].to_numpy(float)
        highs = df["High"].to_numpy(float)
        lows = df["Low"].to_numpy(float)
        closes = df["Close"].to_numpy(float)
        n = len(df)
        c1 = locate_c1(idx, opens, highs, lows, closes, float(rec["Support_Price"]))
        if c1 is None:
            miss_c1 += 1
            rows.append(out)
            continue
        entry = float(rec["Entry_Price"])
        sl = float(rec["SL_Price"])
        trig = float(highs[c1])
        px, xidx, rv = exit_path(idx, entry, sl, trig, opens, highs, lows, closes, n, 2.0, True)
        old_x = pd.Timestamp(rec["Exit_Date"]).normalize()
        new_x = pd.Timestamp(dates.iloc[min(int(xidx), n - 1)]).normalize()
        if new_x != old_x or abs(float(px) - float(rec["Exit_Price"])) > 1e-6:
            n_scratch += 1
        out["Exit_Date"] = new_x
        out["Exit_Price"] = float(px)
        out["Realized_R"] = float(rv)
        out["Outcome"] = "Success" if rv >= 1.5 else ("BE" if rv >= 0 else "Failure")
        rows.append(out)
    df = pd.DataFrame(rows)
    df["Entry_Date"] = pd.to_datetime(df["Entry_Date"])
    df["Exit_Date"] = pd.to_datetime(df["Exit_Date"])
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(CACHE, index=False)
    print(f"[scratch] wrote {CACHE}  n={len(df):,}  exits_changed={n_scratch:,}  miss_c1={miss_c1:,}", flush=True)
    return df


def _chart_rel(index: dict):
    def fn(tid, ticker, entry_dt, net_pnl):
        safe = str(ticker).replace(".NS", "").replace(".BO", "")
        key = (safe, pd.Timestamp(entry_dt).strftime("%Y-%m-%d"))
        name = index.get(key)
        if not name:
            label = "PROFIT" if net_pnl >= 0 else "LOSS"
            name = f"Trade_{tid:04d}_{safe}_{pd.Timestamp(entry_dt).strftime('%Y-%m-%d')}_{label}.png"
        return f"{CHART_REL_PREFIX}/{name}"
    return fn


def sim_statement(
    by_entry: dict,
    paper: dict,
    event_days: list,
    start_dt: pd.Timestamp,
    end_dt: pd.Timestamp,
    ml_n: int,
    closes: CloseCache,
    chart_index: dict,
    write_viz: bool,
) -> dict:
    book = simulate_rr_book(
        by_entry, paper, event_days, start_dt, end_dt, closes,
        chart_rel_fn=_chart_rel(chart_index),
        verbose=False,
        expand_daily=write_viz,
    )
    tag = f"Start_{start_dt.strftime('%Y-%m')}"
    reports = OUT / tag
    plots = PLOTS / tag
    reports.mkdir(parents=True, exist_ok=True)
    plots.mkdir(parents=True, exist_ok=True)

    rows = book["rows"]
    df_daily = book["df_daily"]
    df_stmt = pd.DataFrame(rows)[CSV_COLS]
    df_stmt.to_csv(reports / "Swing_Strategy_Account_Statement.csv", index=False)
    df_daily.to_csv(reports / "Daily_Equity.csv", index=False)

    period = f"{start_dt.strftime('%Y-%m-%d')} to {end_dt.strftime('%Y-%m-%d')}"
    accept_pct = 100.0 * ml_n / max(FEAT_N, 1)
    summary_rows = [
        ("Strategy Rule", RULE, RULE, RULE),
        ("Start vintage", tag, tag, tag),
        ("Initial Capital", f"Rs {CAPITAL:,.2f}", f"Rs {CAPITAL:,.2f}", f"Rs {CAPITAL:,.2f}"),
        ("Risk / Trade", "2% of equity", "2% of equity", "2% of equity"),
        ("Final Portfolio Equity", f"Rs {book['gross_eq']:,.2f}", f"Rs {book['net_z']:,.2f}", f"Rs {book['net_f']:,.2f}"),
        ("Total Net Profit (INR)", f"Rs {book['gross_eq'] - CAPITAL:,.2f}", f"Rs {book['net_z'] - CAPITAL:,.2f}", f"Rs {book['net_f'] - CAPITAL:,.2f}"),
        ("Total Return (%)", f"{book['g_ret']:+.2f}%", f"{book['z_ret']:+.2f}%", f"{book['f_ret']:+.2f}%"),
        ("CAGR (%)", f"{book['g_cagr']:+.2f}%", f"{book['z_cagr']:+.2f}%", f"{book['f_cagr']:+.2f}%"),
        ("Executed Trades Count", f"{book['executed']:,}", f"{book['executed']:,}", f"{book['executed']:,}"),
        ("Win Rate (%)", f"{book['win_rate']:.2f}%", f"{book['win_rate']:.2f}%", f"{book['win_rate']:.2f}%"),
        ("Max Drawdown (%)", f"{book['max_dd']:.2f}%", f"{book['max_dd']:.2f}%", f"{book['max_dd']:.2f}%"),
        ("Total Statutory Taxes Paid", "Rs 0.00", f"Rs {book['total_tax_z']:,.2f}", f"Rs {book['total_tax_f']:,.2f}"),
        ("Backtest Period", period, period, period),
        ("Years (N)", f"{book['years']:.4f}", f"{book['years']:.4f}", f"{book['years']:.4f}"),
        ("Skipped entry days (paper gate)", f"{book['skipped_days']:,}", f"{book['skipped_days']:,}", f"{book['skipped_days']:,}"),
        ("ML accepted / universe", f"{ml_n:,} / {FEAT_N:,}", f"{accept_pct:.1f}%", "Meta_P>=0.48 top16"),
        ("Holding valuation", "MTM qty x close", "MTM qty x close", "MTM qty x close"),
    ]
    _write_excel(rows, summary_rows, reports / "Swing_Strategy_Account_Statement.xlsx")
    if write_viz:
        generate_all_visualizations(df_daily, plots, reports, exp_title=f"Swing_PP scratch {tag}")
    return {
        "start": str(start_dt.date()),
        "end": str(end_dt.date()),
        "years": round(book["years"], 4),
        "Gross_CAGR": round(float(book["g_cagr"]), 2),
        "Zerodha_CAGR": round(float(book["z_cagr"]), 2),
        "FYERS_CAGR": round(float(book["f_cagr"]), 2),
        "Gross_Return": round(float(book["g_ret"]), 2),
        "Zerodha_Return": round(float(book["z_ret"]), 2),
        "Final_Zerodha": round(float(book["net_z"]), 2),
        "Final_Gross": round(float(book["gross_eq"]), 2),
        "DD": round(book["max_dd"], 2),
        "N": book["executed"],
        "Win_Rate": round(book["win_rate"], 2),
        "Taxes_Zerodha": round(book["total_tax_z"], 2),
        "skipped_days": book["skipped_days"],
        "folder": str(reports),
        "open_left": book["open_left"],
        "viz": write_viz,
    }


def _write_master(rows: list[dict], stats: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Start CAGR"
    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    headers = [
        "Experiment ID", "Initial Capital", "Risk Cap / Trade", "Start", "End", "Years",
        "Final Account Equity (Zerodha)", "Net Return (%)", "CAGR (%)",
        "Gross CAGR (%)", "FYERS CAGR (%)", "Executed Trades", "Win Rate (%)",
        "Max Drawdown (%)", "Taxes Paid",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in rows:
        ws.append([
            f"Start_{r['start'][:7]}",
            "Rs 50,000",
            "2% of equity",
            r["start"],
            r["end"],
            r["years"],
            f"Rs {r['Final_Zerodha']:,.2f}",
            f"{r['Zerodha_Return']:+.2f}%",
            f"{r['Zerodha_CAGR']:+.2f}%",
            f"{r['Gross_CAGR']:+.2f}%",
            f"{r['FYERS_CAGR']:+.2f}%",
            r["N"],
            f"{r['Win_Rate']:.2f}%",
            f"{r['DD']:.2f}%",
            f"Rs {r['Taxes_Zerodha']:,.2f}",
        ])
    for i, name in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(36, len(name) + 4))

    ws2 = wb.create_sheet("Performance Summary")
    ws2.append(["Performance Metric", "Gross (BEFORE TAX)", "Net (Zerodha)", "Net (Flat Rs 20)"])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws2.append(["Vintages (month starts)", stats["n_vintages"], stats["n_vintages"], stats["n_vintages"]])
    ws2.append(["Average CAGR (%)", f"{stats['avg_gross']:+.2f}%", f"{stats['avg']:+.2f}%", f"{stats['avg_fyers']:+.2f}%"])
    ws2.append(["Median CAGR (%)", f"{stats['median_gross']:+.2f}%", f"{stats['median']:+.2f}%", f"{stats['median_fyers']:+.2f}%"])
    ws2.append(["Min CAGR (%)", f"{stats['min_gross']:+.2f}%", f"{stats['min']:+.2f}%", f"{stats['min_fyers']:+.2f}%"])
    ws2.append(["Max CAGR (%)", f"{stats['max_gross']:+.2f}%", f"{stats['max']:+.2f}%", f"{stats['max_fyers']:+.2f}%"])
    ws2.append(["Average Max DD (%)", "", f"{stats['avg_dd']:.2f}%", ""])
    ws2.append(["Min Max DD (%)", "", f"{stats['min_dd']:.2f}%", ""])
    ws2.append(["Max Max DD (%)", "", f"{stats['max_dd']:.2f}%", ""])
    ws2.append(["Capital", "Rs 50,000", "Rs 50,000", "Rs 50,000"])
    ws2.append(["Risk / Trade", "2% of equity", "2% of equity", "2% of equity"])
    ws2.append(["Rule", RULE, RULE, RULE])
    for i in range(1, 5):
        ws2.column_dimensions[get_column_letter(i)].width = 42
    wb.save(MASTER)


def _plot_vintage_heatmap(df: pd.DataFrame) -> None:
    d = df.copy()
    d["year"] = pd.to_datetime(d["start"]).dt.year
    d["month"] = pd.to_datetime(d["start"]).dt.month
    pivot = d.pivot(index="year", columns="month", values="Zerodha_CAGR")
    month_names = {1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "May", 6: "Jun",
                   7: "Jul", 8: "Aug", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dec"}
    pivot = pivot.rename(columns=month_names)
    PLOTS.mkdir(parents=True, exist_ok=True)
    plt.figure(figsize=(12, max(6, len(pivot) * 0.45)))
    sns.set_theme(style="white")
    ax = sns.heatmap(
        pivot, annot=True, fmt="+.1f",
        cmap=sns.diverging_palette(10, 130, as_cmap=True),
        center=float(d["Zerodha_CAGR"].median()),
        cbar_kws={"label": "Net Zerodha CAGR (%)"},
        linewidths=0.8, linecolor="white",
        annot_kws={"size": 8, "weight": "bold"},
    )
    ax.set_title("Swing_PP scratch — Net Zerodha CAGR by month-start vintage", fontsize=13, fontweight="bold", pad=12)
    ax.set_xlabel("Start month")
    ax.set_ylabel("Start year")
    plt.tight_layout()
    plt.savefig(HEATMAP, dpi=200)
    plt.close()


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    base = load_pred_list()
    ml = rewrite_scratch(base)
    paper = exits_by_day(ml, 1.5)
    by_entry = defaultdict(list)
    for rec in ml.to_dict("records"):
        if pd.isna(rec.get("Exit_Date")):
            continue
        by_entry[pd.Timestamp(rec["Entry_Date"]).normalize()].append(rec)
    for day, cands in by_entry.items():
        cands.sort(key=lambda x: float(x.get("Meta_P", 0.0)), reverse=True)

    first_entry = min(by_entry.keys())
    end_dt = max(ml["Entry_Date"].max(), pd.to_datetime(ml["Exit_Date"]).max()).normalize()
    extra_x = []
    for recs in by_entry.values():
        for r in recs:
            xd = pd.Timestamp(r["Exit_Date"])
            if pd.notna(xd):
                extra_x.append(xd.normalize())
    event_days = sorted(set(by_entry.keys()) | set(paper.keys()) | set(extra_x))

    closes = CloseCache()
    tickers = {str(r.get("Ticker")) for recs in by_entry.values() for r in recs}
    print(f"Loading {len(tickers)} daily close series for MTM ...", flush=True)
    closes.preload(tickers)
    CHARTS.mkdir(parents=True, exist_ok=True)
    chart_index = index_trade_charts(CHARTS)
    print(f"trade-chart index {len(chart_index):,} files", flush=True)

    starts = pd.date_range("2005-01-01", end_dt, freq="MS")
    print(
        f"prediction list {len(ml):,}  first entry {first_entry.date()}  "
        f"end {end_dt.date()}  vintages {len(starts)}",
        flush=True,
    )
    OUT.mkdir(parents=True, exist_ok=True)
    PLOTS.mkdir(parents=True, exist_ok=True)

    rows = []
    for i, start in enumerate(starts, 1):
        start_dt = pd.Timestamp(start).normalize()
        write_viz = start_dt.month == 1 or start_dt == pd.Timestamp("2005-01-01")
        rec = sim_statement(
            by_entry, paper, event_days, start_dt, end_dt, len(ml),
            closes, chart_index, write_viz,
        )
        rows.append(rec)
        if i == 1 or start_dt.month == 1 or i == len(starts) or i % 12 == 0:
            print(
                f"  {rec['start']}  Zerodha CAGR {rec['Zerodha_CAGR']:+7.2f}%  "
                f"DD {rec['DD']:5.1f}%  n={rec['N']:4d}  {rec['years']:.2f}y  "
                f"final Rs {rec['Final_Zerodha']:,.0f}",
                flush=True,
            )

    df = pd.DataFrame(rows)
    z = df["Zerodha_CAGR"]
    g = df["Gross_CAGR"]
    f = df["FYERS_CAGR"]
    long = df[df["years"] >= 3]
    stats = {
        "n_vintages": int(len(df)),
        "first_entry": str(first_entry.date()),
        "end": str(end_dt.date()),
        "avg": round(float(z.mean()), 2),
        "median": round(float(z.median()), 2),
        "min": round(float(z.min()), 2),
        "max": round(float(z.max()), 2),
        "avg_gross": round(float(g.mean()), 2),
        "median_gross": round(float(g.median()), 2),
        "min_gross": round(float(g.min()), 2),
        "max_gross": round(float(g.max()), 2),
        "avg_fyers": round(float(f.mean()), 2),
        "median_fyers": round(float(f.median()), 2),
        "min_fyers": round(float(f.min()), 2),
        "max_fyers": round(float(f.max()), 2),
        "avg_dd": round(float(df["DD"].mean()), 2),
        "min_dd": round(float(df["DD"].min()), 2),
        "max_dd": round(float(df["DD"].max()), 2),
        "min_start": df.loc[z.idxmin(), "start"],
        "max_start": df.loc[z.idxmax(), "start"],
        "avg_ge3y": round(float(long["Zerodha_CAGR"].mean()), 2) if len(long) else None,
        "min_ge3y": round(float(long["Zerodha_CAGR"].min()), 2) if len(long) else None,
        "max_ge3y": round(float(long["Zerodha_CAGR"].max()), 2) if len(long) else None,
        "n_ge3y": int(len(long)),
    }
    df.to_csv(CSV, index=False)
    _write_master(rows, stats)
    _plot_vintage_heatmap(df)
    LOG.write_text(
        json.dumps({"rule": RULE, "capital": CAPITAL, "stats": stats, "rows": rows}, indent=2, default=str),
        encoding="utf-8",
    )
    print("\n===== Net Zerodha CAGR (all month starts) =====", flush=True)
    print(f"n={stats['n_vintages']}  avg {stats['avg']:+.2f}%  min {stats['min']:+.2f}% ({stats['min_start']})  max {stats['max']:+.2f}% ({stats['max_start']})", flush=True)
    print(f">=3y n={stats['n_ge3y']}  avg {stats['avg_ge3y']:+.2f}%  min {stats['min_ge3y']:+.2f}%  max {stats['max_ge3y']:+.2f}%", flush=True)
    print(f"DD avg {stats['avg_dd']:.1f}%  min {stats['min_dd']:.1f}%  max {stats['max_dd']:.1f}%", flush=True)
    print(f"wrote {CSV}", flush=True)
    print(f"wrote {MASTER}", flush=True)
    print(f"wrote {HEATMAP}", flush=True)


if __name__ == "__main__":
    main()
