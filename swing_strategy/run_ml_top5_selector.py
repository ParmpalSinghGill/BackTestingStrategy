"""
Walk-forward ML/DL selector: among same-day C1/C2 setups, score with
models trained only on the past, then take the predicted top N.

Zero lookahead in features (bars strictly before Entry_Date, plus today's open).
Objective: Net Zerodha CAGR 40%+ on 50k/500 and/or 100k/500, 1:2 book.
"""

from __future__ import annotations

import json
import sys
import warnings
from collections import defaultdict
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

warnings.filterwarnings("ignore")

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from src.analysis.indian_brokerage_calculator import calculate_indian_trade_charges
from swing_strategy.tiered_liquidity_strategy_engine import DATA_DAILY_DIR
from swing_strategy.visualizer import generate_all_visualizations

SETUP_CSV = BASE_DIR / "Reports" / "TopN_Oracle_Select" / "All_Setups_RR2.csv"
FEAT_PATH = BASE_DIR / "Reports" / "ML_Top5_Selector" / "Features_RR2.parquet"
OUT_BASE = BASE_DIR / "Reports" / "ML_Top5_Selector"
PLOT_BASE = BASE_DIR / "Plots" / "ML_Top5_Selector"
LEADERBOARD = OUT_BASE / "Leaderboard.csv"

FEATURE_COLS = [
    "TF_Rank", "Nifty_Rank", "risk_pct", "dist_support_pct",
    "ret_5", "ret_10", "ret_20", "ret_60",
    "sma20_dist", "sma50_dist", "sma200_dist",
    "atr14_pct", "vol20", "rsi14", "vol_ratio20",
    "dist_high20", "dist_low20", "gap_pct",
    "dow", "month",
    "p1_color", "p1_body", "p1_range",
    "p2_color", "p2_body",
    "p3_color", "p3_body",
    "n_cands", "rel_ret20", "rel_vol", "rel_risk", "rel_tf", "rel_nifty",
]


def _rsi(close: np.ndarray, n: int = 14) -> np.ndarray:
    delta = np.diff(close, prepend=close[0])
    gain = np.where(delta > 0, delta, 0.0)
    loss = np.where(delta < 0, -delta, 0.0)
    out = np.full(len(close), 50.0)
    if len(close) <= n:
        return out
    ag = gain[1 : n + 1].mean()
    al = loss[1 : n + 1].mean()
    for i in range(n, len(close)):
        ag = (ag * (n - 1) + gain[i]) / n
        al = (al * (n - 1) + loss[i]) / n
        out[i] = 100.0 if al == 0 else 100.0 - 100.0 / (1.0 + ag / al)
    return out


def _sma(x: np.ndarray, n: int) -> np.ndarray:
    if len(x) < n:
        return np.full(len(x), np.nan)
    c = np.cumsum(x)
    out = np.full(len(x), np.nan)
    out[n - 1 :] = (c[n - 1 :] - np.concatenate([[0], c[:-n]])) / n
    return out


def feat_worker(payload: tuple) -> list[dict]:
    symbol, recs = payload
    path = DATA_DAILY_DIR / f"{symbol}_1d.csv"
    if not path.exists() or not recs:
        return []
    try:
        df = pd.read_csv(path)
    except Exception:
        return []
    if "Date" not in df.columns or len(df) < 80:
        return []
    df["Date"] = pd.to_datetime(df["Date"])
    df = df.sort_values("Date").reset_index(drop=True)
    opens = df["Open"].to_numpy(float)
    highs = df["High"].to_numpy(float)
    lows = df["Low"].to_numpy(float)
    closes = df["Close"].to_numpy(float)
    vols = df["Volume"].to_numpy(float) if "Volume" in df.columns else np.ones(len(df))
    dates = df["Date"]
    date_to_i = {pd.Timestamp(d).normalize(): i for i, d in enumerate(dates)}
    n = len(df)

    logc = np.log(np.clip(closes, 1e-6, None))
    rets = np.diff(logc, prepend=logc[0])
    sma20 = _sma(closes, 20)
    sma50 = _sma(closes, 50)
    sma200 = _sma(closes, 200)
    tr = np.maximum(highs - lows, np.maximum(np.abs(highs - np.roll(closes, 1)), np.abs(lows - np.roll(closes, 1))))
    tr[0] = highs[0] - lows[0]
    atr14 = _sma(tr, 14)
    vol20 = pd.Series(rets).rolling(20, min_periods=10).std().to_numpy()
    rsi = _rsi(closes, 14)
    vsma = _sma(vols, 20)
    hh20 = pd.Series(highs).rolling(20, min_periods=5).max().to_numpy()
    ll20 = pd.Series(lows).rolling(20, min_periods=5).min().to_numpy()

    out = []
    for rec in recs:
        edt = pd.Timestamp(rec["Entry_Date"]).normalize()
        i = date_to_i.get(edt)
        if i is None or i < 65:
            continue
        p = i - 1
        c = float(closes[p])
        if c <= 0:
            continue
        o_today = float(opens[i])

        def ret_n(k):
            j = p - k
            if j < 0 or closes[j] <= 0:
                return 0.0
            return float(closes[p] / closes[j] - 1.0)

        def candle(j):
            if j < 0:
                return 0.0, 0.0, 0.0
            o, h, l, cl = opens[j], highs[j], lows[j], closes[j]
            rng = max(h - l, 1e-6)
            color = 1.0 if cl > o else (-1.0 if cl < o else 0.0)
            body = abs(cl - o) / rng
            return color, body, rng / max(cl, 1e-6)

        c1c, c1b, c1r = candle(p)
        c2c, c2b, _ = candle(p - 1)
        c3c, c3b, _ = candle(p - 2)
        entry = float(rec["Entry_Price"])
        sl = float(rec["SL_Price"])
        sup = float(rec["Support_Price"])
        risk = entry - sl
        row = dict(rec)
        row.update({
            "risk_pct": risk / entry if entry else 0.0,
            "dist_support_pct": (entry - sup) / sup if sup else 0.0,
            "ret_5": ret_n(5),
            "ret_10": ret_n(10),
            "ret_20": ret_n(20),
            "ret_60": ret_n(60),
            "sma20_dist": (c / sma20[p] - 1.0) if sma20[p] == sma20[p] and sma20[p] else 0.0,
            "sma50_dist": (c / sma50[p] - 1.0) if sma50[p] == sma50[p] and sma50[p] else 0.0,
            "sma200_dist": (c / sma200[p] - 1.0) if sma200[p] == sma200[p] and sma200[p] else 0.0,
            "atr14_pct": float(atr14[p] / c) if atr14[p] == atr14[p] else 0.0,
            "vol20": float(vol20[p]) if vol20[p] == vol20[p] else 0.0,
            "rsi14": float(rsi[p]) if rsi[p] == rsi[p] else 50.0,
            "vol_ratio20": float(vols[p] / vsma[p]) if vsma[p] == vsma[p] and vsma[p] else 1.0,
            "dist_high20": (c / hh20[p] - 1.0) if hh20[p] == hh20[p] and hh20[p] else 0.0,
            "dist_low20": (c / ll20[p] - 1.0) if ll20[p] == ll20[p] and ll20[p] else 0.0,
            "gap_pct": o_today / c - 1.0,
            "dow": float(edt.dayofweek),
            "month": float(edt.month),
            "p1_color": c1c, "p1_body": c1b, "p1_range": c1r,
            "p2_color": c2c, "p2_body": c2b,
            "p3_color": c3c, "p3_body": c3b,
        })
        out.append(row)
    return out


def build_features() -> pd.DataFrame:
    if FEAT_PATH.exists():
        print(f"Loading features {FEAT_PATH}", flush=True)
        return pd.read_parquet(FEAT_PATH)
    raw = pd.read_csv(SETUP_CSV)
    by = defaultdict(list)
    for rec in raw.to_dict("records"):
        by[rec["Ticker"]].append(rec)
    jobs = list(by.items())
    rows: list[dict] = []
    print(f"Building features for {len(jobs):,} tickers ...", flush=True)
    done = 0
    with ProcessPoolExecutor(max_workers=8) as ex:
        futs = {ex.submit(feat_worker, j): j[0] for j in jobs}
        for fut in as_completed(futs):
            rows.extend(fut.result())
            done += 1
            if done % 300 == 0 or done == len(jobs):
                print(f"  • {done:,}/{len(jobs):,} tickers | rows {len(rows):,}", flush=True)
    df = pd.DataFrame(rows)
    df["Entry_Date"] = pd.to_datetime(df["Entry_Date"])
    df["Exit_Date"] = pd.to_datetime(df["Exit_Date"])
    df["year"] = df["Entry_Date"].dt.year
    df["day_rank"] = df.groupby("Entry_Date")["Realized_R"].rank(ascending=False, method="first")
    df["y_top5"] = (df["day_rank"] <= 5).astype(int)
    df["y_win"] = (df["Outcome"] == "Success").astype(int)
    df["y_r"] = df["Realized_R"].clip(-2.0, 4.0)
    g = df.groupby("Entry_Date")
    df["n_cands"] = g["Ticker"].transform("size")
    df["rel_ret20"] = g["ret_20"].rank(pct=True)
    df["rel_vol"] = g["vol_ratio20"].rank(pct=True)
    df["rel_risk"] = g["risk_pct"].rank(pct=True)
    df["rel_tf"] = g["TF_Rank"].rank(pct=True)
    df["rel_nifty"] = g["Nifty_Rank"].rank(pct=True)
    for c in FEATURE_COLS:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    FEAT_PATH.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(FEAT_PATH, index=False)
    print(f"Saved {FEAT_PATH} rows={len(df):,} top5_rate={df.y_top5.mean():.3f} win={df.y_win.mean():.3f}", flush=True)
    return df


def _X(df: pd.DataFrame) -> np.ndarray:
    return df[FEATURE_COLS].to_numpy(dtype=np.float32)


def walk_forward(df: pd.DataFrame, kind: str) -> pd.DataFrame:
    """kind: lgbm_rank | lgbm_cls | lgbm_reg | xgb_cls | xgb_reg | hgb_cls | mlp_cls"""
    parts = []
    years = sorted(df["year"].unique())
    for y in years:
        train = df[df["year"] < y]
        test = df[df["year"] == y].copy()
        if test.empty:
            continue
        if len(train) < 800:
            test["ML_Score"] = test["TF_Rank"] * 10 + test["Nifty_Rank"]
            parts.append(test)
            continue
        Xtr, Xte = _X(train), _X(test)
        score = None
        if kind == "lgbm_rank":
            import lightgbm as lgb
            tr = train.sort_values("Entry_Date")
            groups = tr.groupby("Entry_Date", sort=True).size().tolist()
            rel = np.clip(np.rint(tr["y_r"].to_numpy() + 2.0), 0, 6).astype(int)
            dtrain = lgb.Dataset(_X(tr), label=rel, group=groups, feature_name=FEATURE_COLS)
            params = {
                "objective": "lambdarank",
                "metric": "ndcg",
                "ndcg_eval_at": [5],
                "learning_rate": 0.05,
                "num_leaves": 48,
                "min_data_in_leaf": 80,
                "feature_fraction": 0.8,
                "verbosity": -1,
            }
            model = lgb.train(params, dtrain, num_boost_round=250)
            score = model.predict(Xte)
        elif kind == "lgbm_cls":
            import lightgbm as lgb
            dtrain = lgb.Dataset(Xtr, label=train["y_win"].to_numpy(), feature_name=FEATURE_COLS)
            params = {
                "objective": "binary",
                "metric": "auc",
                "learning_rate": 0.05,
                "num_leaves": 48,
                "min_data_in_leaf": 80,
                "feature_fraction": 0.8,
                "verbosity": -1,
            }
            model = lgb.train(params, dtrain, num_boost_round=300)
            score = model.predict(Xte)
        elif kind == "lgbm_reg":
            import lightgbm as lgb
            dtrain = lgb.Dataset(Xtr, label=train["y_r"].to_numpy(), feature_name=FEATURE_COLS)
            params = {
                "objective": "regression",
                "metric": "rmse",
                "learning_rate": 0.05,
                "num_leaves": 48,
                "min_data_in_leaf": 80,
                "verbosity": -1,
            }
            model = lgb.train(params, dtrain, num_boost_round=300)
            score = model.predict(Xte)
        elif kind == "lgbm_top5":
            import lightgbm as lgb
            dtrain = lgb.Dataset(Xtr, label=train["y_top5"].to_numpy(), feature_name=FEATURE_COLS)
            pos = max(int(train["y_top5"].sum()), 1)
            neg = max(len(train) - pos, 1)
            params = {
                "objective": "binary",
                "metric": "auc",
                "learning_rate": 0.05,
                "num_leaves": 63,
                "min_data_in_leaf": 60,
                "scale_pos_weight": neg / pos,
                "verbosity": -1,
            }
            model = lgb.train(params, dtrain, num_boost_round=350)
            score = model.predict(Xte)
        elif kind == "xgb_cls":
            from xgboost import XGBClassifier
            clf = XGBClassifier(
                n_estimators=250, max_depth=5, learning_rate=0.05,
                subsample=0.8, colsample_bytree=0.8, n_jobs=4,
                eval_metric="logloss", tree_method="hist",
            )
            clf.fit(Xtr, train["y_win"].to_numpy())
            score = clf.predict_proba(Xte)[:, 1]
        elif kind == "xgb_reg":
            from xgboost import XGBRegressor
            reg = XGBRegressor(
                n_estimators=250, max_depth=5, learning_rate=0.05,
                subsample=0.8, colsample_bytree=0.8, n_jobs=4, tree_method="hist",
            )
            reg.fit(Xtr, train["y_r"].to_numpy())
            score = reg.predict(Xte)
        elif kind == "hgb_cls":
            from sklearn.ensemble import HistGradientBoostingClassifier
            clf = HistGradientBoostingClassifier(max_depth=6, learning_rate=0.06, max_iter=200)
            clf.fit(Xtr, train["y_win"].to_numpy())
            score = clf.predict_proba(Xte)[:, 1]
        elif kind == "mlp_cls":
            from sklearn.neural_network import MLPClassifier
            from sklearn.preprocessing import StandardScaler
            sc = StandardScaler()
            clf = MLPClassifier(
                hidden_layer_sizes=(64, 32), activation="relu", max_iter=40,
                random_state=42, early_stopping=True, n_iter_no_change=5,
            )
            clf.fit(sc.fit_transform(Xtr), train["y_win"].to_numpy())
            score = clf.predict_proba(sc.transform(Xte))[:, 1]
        else:
            raise ValueError(kind)
        test["ML_Score"] = score
        parts.append(test)
        print(f"    {kind} year {y}: train {len(train):,} test {len(test):,}", flush=True)
    return pd.concat(parts, ignore_index=True)


def _metrics(initial: float, final_eq: float, years: float) -> tuple[float, float]:
    ret = ((final_eq - initial) / initial) * 100.0
    if final_eq <= 0:
        return round(ret, 2), -100.0
    cagr = (((final_eq / initial) ** (1.0 / max(years, 0.01))) - 1.0) * 100.0
    return round(ret, 2), round(cagr, 2)


def run_portfolio(df_trades: pd.DataFrame, capital: float, risk: float, exp_name: str, write_files: bool) -> dict:
    reports_dir = OUT_BASE / exp_name
    plots_dir = PLOT_BASE / exp_name
    if write_files:
        reports_dir.mkdir(parents=True, exist_ok=True)
        plots_dir.mkdir(parents=True, exist_ok=True)

    df = df_trades.copy()
    df["Entry_Date"] = pd.to_datetime(df["Entry_Date"])
    by_entry: dict = defaultdict(list)
    for row in df.to_dict("records"):
        by_entry[row["Entry_Date"]].append(row)

    min_dt = pd.Timestamp("2010-01-01")
    max_dt = max(df["Entry_Date"].max(), pd.to_datetime(df["Exit_Date"]).max())
    cash = capital
    peak = capital
    max_dd = 0.0
    tx_id = 2
    trade_id = 1
    open_pos: dict[int, dict] = {}
    executed = 0
    total_tax_z = total_tax_f = total_gross = 0.0
    rows = [{
        "Transaction_ID": 1, "Trade_ID": 0, "Type": "DEPOSIT", "Date": "2010-01-01",
        "Ticker": "N/A", "Liquidity_Source": "N/A", "Support_Price": 0.0, "Quantity": 0,
        "Price": 0.0, "Total_Spend": 0.0, "Gross_PnL": 0.0, "Statutory_Taxes": 0.0,
        "Net_PnL": 0.0, "Return_Pct": 0.0, "Cash_Balance": cash,
        "Active_Position_Count": 0, "Holding_Equity_Value": 0.0,
        "Total_Portfolio_Value": cash, "Target_RR_Mode": "1:2",
        "Outcome": "DEPOSIT", "Chart_PNG_URI": "N/A",
    }]
    daily = []
    event_days = sorted(set(by_entry.keys()) | {
        pd.Timestamp(r["Exit_Date"]) for recs in by_entry.values() for r in recs
    })
    if write_files:
        day_iter = list(pd.date_range(min_dt, max_dt, freq="D"))
    else:
        day_iter = event_days

    for day in day_iter:
        day_start = cash
        day_pnl = 0.0
        avail = cash
        cands = by_entry.get(day, [])
        for cand in cands:
            entry_p = float(cand["Entry_Price"])
            sl_p = float(cand["SL_Price"])
            rsk = entry_p - sl_p
            if rsk <= 0.05 or rsk > risk:
                continue
            qty = min(int(risk // rsk), int(avail // entry_p))
            if qty < 1:
                continue
            spend = round(entry_p * qty, 2)
            if spend > avail:
                continue
            cash = round(cash - spend, 2)
            avail = round(avail - spend, 2)
            executed += 1
            open_pos[trade_id] = {
                "Ticker": cand["Ticker"],
                "Liquidity_Source": cand.get("Liquidity_Type", "Weekly"),
                "Support_Price": float(cand["Support_Price"]),
                "Entry_Price": entry_p,
                "Exit_Date": pd.Timestamp(cand["Exit_Date"]),
                "Exit_Price": float(cand["Exit_Price"]),
                "Quantity": qty,
                "Total_Spend": spend,
            }
            holding = sum(p["Total_Spend"] for p in open_pos.values())
            rows.append({
                "Transaction_ID": tx_id, "Trade_ID": trade_id, "Type": "BUY (ENTRY)",
                "Date": day.strftime("%Y-%m-%d"), "Ticker": cand["Ticker"],
                "Liquidity_Source": cand.get("Liquidity_Type", "Weekly"),
                "Support_Price": float(cand["Support_Price"]),
                "Quantity": qty, "Price": entry_p, "Total_Spend": spend,
                "Gross_PnL": 0.0, "Statutory_Taxes": 0.0, "Net_PnL": 0.0, "Return_Pct": 0.0,
                "Cash_Balance": cash, "Active_Position_Count": len(open_pos),
                "Holding_Equity_Value": holding, "Total_Portfolio_Value": cash + holding,
                "Target_RR_Mode": "1:2", "Outcome": "OPEN", "Chart_PNG_URI": "N/A",
            })
            tx_id += 1
            trade_id += 1

        for tid, pos in list(open_pos.items()):
            if pos["Exit_Date"] > day:
                continue
            qty = pos["Quantity"]
            entry_p, exit_p = pos["Entry_Price"], pos["Exit_Price"]
            spend = pos["Total_Spend"]
            gross = round((exit_p - entry_p) * qty, 2)
            ch_z = calculate_indian_trade_charges(entry_p, exit_p, qty, 0.0)
            ch_f = calculate_indian_trade_charges(entry_p, exit_p, qty, 20.0)
            tax_z = round(ch_z["total_charges"], 2)
            tax_f = round(ch_f["total_charges"], 2)
            net = round(gross - tax_z, 2)
            total_gross += gross
            total_tax_z += tax_z
            total_tax_f += tax_f
            day_pnl += net
            cash = round(cash + spend + gross - tax_z, 2)
            del open_pos[tid]
            holding = sum(p["Total_Spend"] for p in open_pos.values())
            rows.append({
                "Transaction_ID": tx_id, "Trade_ID": tid, "Type": "SELL (EXIT)",
                "Date": pos["Exit_Date"].strftime("%Y-%m-%d"), "Ticker": pos["Ticker"],
                "Liquidity_Source": pos["Liquidity_Source"], "Support_Price": pos["Support_Price"],
                "Quantity": qty, "Price": exit_p, "Total_Spend": spend,
                "Gross_PnL": gross, "Statutory_Taxes": tax_z, "Net_PnL": net,
                "Return_Pct": round((net / spend) * 100, 2) if spend else 0.0,
                "Cash_Balance": cash, "Active_Position_Count": len(open_pos),
                "Holding_Equity_Value": holding, "Total_Portfolio_Value": cash + holding,
                "Target_RR_Mode": "1:2",
                "Outcome": "Success" if net >= 0 else "Failure",
                "Chart_PNG_URI": "N/A",
            })
            tx_id += 1

        port = cash + sum(p["Total_Spend"] for p in open_pos.values())
        peak = max(peak, port)
        max_dd = max(max_dd, ((peak - port) / peak * 100) if peak > 0 else 0)
        daily.append({
            "Date": day.strftime("%Y-%m-%d"),
            "Balance": cash,
            "Active_Positions": len(open_pos),
            "Daily_PnL": day_pnl,
            "Daily_Return_Pct": ((cash - day_start) / day_start * 100) if day_start else 0,
        })

    sells = [r for r in rows if r["Type"] == "SELL (EXIT)"]
    wins = sum(1 for r in sells if r["Net_PnL"] >= 0)
    win_rate = (wins / len(sells) * 100) if sells else 0.0
    years = max((max_dt - min_dt).days / 365.25, 0.01)
    gross_eq = round(capital + total_gross, 2)
    net_z = cash
    net_f = round(net_z - (total_tax_f - total_tax_z), 2)
    g_ret, g_cagr = _metrics(capital, gross_eq, years)
    z_ret, z_cagr = _metrics(capital, net_z, years)
    f_ret, f_cagr = _metrics(capital, net_f, years)

    if write_files:
        df_stmt = pd.DataFrame(rows)
        df_daily = pd.DataFrame(daily)
        df_stmt.to_csv(reports_dir / "Swing_Strategy_Account_Statement.csv", index=False)
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Account Statement"
        headers = list(df_stmt.columns)
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center")
        for _, r in df_stmt.iterrows():
            ws.append(list(r))
        ws_sum = wb.create_sheet("Performance Summary")
        ws_sum.append(["Performance Metric", "Gross (BEFORE TAX)", "Net (Zerodha)", "Net (Flat Rs 20)"])
        period = f"{min_dt.strftime('%Y-%m-%d')} to {max_dt.strftime('%Y-%m-%d')}"
        for row in [
            ("Strategy Rule", "Walk-forward ML top-N of C1/C2 setups (no lookahead)",) * 4,
            ("Initial Capital", f"Rs {capital:,.2f}", f"Rs {capital:,.2f}", f"Rs {capital:,.2f}"),
            ("Fixed Risk Cap / Trade", f"Rs {risk:,.2f}", f"Rs {risk:,.2f}", f"Rs {risk:,.2f}"),
            ("Final Portfolio Equity", f"Rs {gross_eq:,.2f}", f"Rs {net_z:,.2f}", f"Rs {net_f:,.2f}"),
            ("Total Net Profit (INR)", f"Rs {gross_eq - capital:,.2f}", f"Rs {net_z - capital:,.2f}", f"Rs {net_f - capital:,.2f}"),
            ("Total Return (%)", f"{g_ret:+.2f}%", f"{z_ret:+.2f}%", f"{f_ret:+.2f}%"),
            ("CAGR (%)", f"{g_cagr:+.2f}%", f"{z_cagr:+.2f}%", f"{f_cagr:+.2f}%"),
            ("Executed Trades Count", f"{executed:,}", f"{executed:,}", f"{executed:,}"),
            ("Win Rate (%)", f"{win_rate:.2f}%", f"{win_rate:.2f}%", f"{win_rate:.2f}%"),
            ("Max Drawdown (%)", f"{max_dd:.2f}%", f"{max_dd:.2f}%", f"{max_dd:.2f}%"),
            ("Total Statutory Taxes Paid", "Rs 0.00", f"Rs {total_tax_z:,.2f}", f"Rs {total_tax_f:,.2f}"),
            ("Backtest Period", period, period, period),
        ]:
            ws_sum.append(list(row))
        wb.save(reports_dir / "Swing_Strategy_Account_Statement.xlsx")
        generate_all_visualizations(df_daily, plots_dir, reports_dir, exp_title=exp_name)

    return {
        "Experiment": exp_name,
        "Capital": capital,
        "Final_Zerodha": round(net_z, 2),
        "Net_Return_Pct": z_ret,
        "CAGR_Pct": z_cagr,
        "Gross_CAGR_Pct": g_cagr,
        "Fyers_CAGR_Pct": f_cagr,
        "Executed": executed,
        "Win_Rate_Pct": round(win_rate, 2),
        "Max_DD_Pct": round(max_dd, 2),
    }


def select_day(scored: pd.DataFrame, top_n: int, min_score: float | None) -> pd.DataFrame:
    keep = []
    for _, g in scored.groupby("Entry_Date", sort=False):
        g = g.sort_values("ML_Score", ascending=False)
        if min_score is not None:
            g = g[g["ML_Score"] >= min_score]
        keep.append(g.head(top_n))
    if not keep:
        return scored.iloc[0:0]
    return pd.concat(keep, ignore_index=True)


def eval_config(scored: pd.DataFrame, kind: str, top_n: int, min_score: float | None, write_best: bool) -> list[dict]:
    picked = select_day(scored, top_n, min_score)
    tag = f"{kind}_top{top_n}" + (f"_p{min_score:.2f}" if min_score is not None else "")
    print(f"  Eval {tag}: candidates {len(picked):,}", flush=True)
    rows = []
    for cap, cname in ((50_000.0, "50k"), (100_000.0, "100k")):
        exp = f"{tag}_Cap{cname}_Risk500"
        res = run_portfolio(picked, cap, 500.0, exp, write_files=write_best)
        res["Model"] = kind
        res["Top_N"] = top_n
        res["Min_Score"] = min_score if min_score is not None else ""
        print(
            f"    {exp}: Rs {res['Final_Zerodha']:,.0f} | CAGR {res['CAGR_Pct']:+.2f}% | "
            f"WR {res['Win_Rate_Pct']:.1f}% | n={res['Executed']}",
            flush=True,
        )
        rows.append(res)
    return rows


def main():
    OUT_BASE.mkdir(parents=True, exist_ok=True)
    df = build_features()
    print(
        f"Dataset {len(df):,} | days {df.Entry_Date.nunique():,} | "
        f"top5 {df.y_top5.mean()*100:.1f}% | win {df.y_win.mean()*100:.1f}%",
        flush=True,
    )

    kinds = [
        "lgbm_rank",
        "lgbm_cls",
        "lgbm_reg",
        "lgbm_top5",
        "xgb_cls",
        "xgb_reg",
        "hgb_cls",
        "mlp_cls",
    ]
    scored_by: dict[str, pd.DataFrame] = {}
    for kind in kinds:
        print(f"\nWalk-forward {kind}", flush=True)
        try:
            scored_by[kind] = walk_forward(df, kind)
            scored_by[kind].to_parquet(OUT_BASE / f"Scored_{kind}.parquet", index=False)
        except Exception as exc:
            print(f"  FAIL {kind}: {exc}", flush=True)

    configs = []
    for kind in scored_by:
        configs.append((kind, 5, None))
        configs.append((kind, 3, None))
        configs.append((kind, 1, None))
        if kind in ("lgbm_cls", "lgbm_top5", "xgb_cls", "hgb_cls", "mlp_cls"):
            for thr in (0.50, 0.55, 0.60, 0.65):
                configs.append((kind, 5, thr))
                configs.append((kind, 3, thr))
                configs.append((kind, 1, thr))

    board: list[dict] = []
    best_cagr = -999.0
    best_row = None
    for kind, top_n, thr in configs:
        rows = eval_config(scored_by[kind], kind, top_n, thr, write_best=False)
        board.extend(rows)
        for r in rows:
            if r["CAGR_Pct"] > best_cagr:
                best_cagr = r["CAGR_Pct"]
                best_row = r
        pd.DataFrame(board).to_csv(LEADERBOARD, index=False)

    if best_row:
        print(f"\nBEST so far: {best_row['Experiment']} CAGR {best_cagr:+.2f}%", flush=True)
        picked = select_day(
            scored_by[best_row["Model"]],
            int(best_row["Top_N"]),
            None if best_row["Min_Score"] == "" else float(best_row["Min_Score"]),
        )
        run_portfolio(
            picked, float(best_row["Capital"]), 500.0, best_row["Experiment"], write_files=True,
        )

    print("\n" + pd.DataFrame(board).sort_values("CAGR_Pct", ascending=False).head(20).to_string(index=False), flush=True)
    print(f"\nSaved {LEADERBOARD}", flush=True)
    (OUT_BASE / "best.json").write_text(json.dumps(best_row, indent=2), encoding="utf-8")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
