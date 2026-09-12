"""Year-start vintages of the Swing_PP RR stack (M48).

Fresh Rs 50,000 on 1 Jan of each year 2010..2026, hold through last data date.
Same rule: after +1R, P6>=0.85->1:6 else P5>=0.85->1:5 else P3>=0.80->1:3 else 1:2.
Paper-gate seeded from 1:2 paper exits before the start date.

21-col statements use MTM holding (qty x close) per Guide/Account_Statement_guide.md.
Chart_PNG_URI hyperlinks reuse Plots/SwingPP_RR/trade_charts when the same ticker+entry exists.
Does not write Swing_low / Swing_Live.
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

import pandas as pd

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from swing_strategy.run_2pct_50_25_hunt import CAPITAL
from swing_strategy.run_pred_paper_gate import exits_by_day
from swing_strategy.run_swing_pp_rr_statement import FEAT_N, _scored_fill, simulate_rr_book
from swing_strategy.run_swing_pp_statement import CSV_COLS, _write_excel
from swing_strategy.statement_mtm import CloseCache, index_trade_charts
from swing_strategy.visualizer import generate_all_visualizations

OUT = BASE / "Reports" / "SwingPP_RR_StartYears"
PLOTS = BASE / "Plots" / "SwingPP_RR_StartYears"
CHARTS_2010 = BASE / "Plots" / "SwingPP_RR" / "trade_charts"
MASTER = OUT / "Master_StartYear_Comparison.xlsx"
CSV = OUT / "StartYear_CAGR.csv"
LOG = OUT / "StartYear_CAGR.json"
RULE = (
    "Swing_PP RR stack | after +1R: P6>=0.85->1:6 else P5>=0.85->1:5 "
    "else P3>=0.80->1:3 else 1:2 | fair BE | 2% equity | paper-gate last 50 fail%>=74%"
)
CHART_REL_PREFIX = "../../../Plots/SwingPP_RR/trade_charts"


def _year_chart_rel(index: dict):
    def fn(tid, ticker, entry_dt, net_pnl):
        safe = str(ticker).replace(".NS", "").replace(".BO", "")
        key = (safe, pd.Timestamp(entry_dt).strftime("%Y-%m-%d"))
        name = index.get(key)
        if name:
            return f"{CHART_REL_PREFIX}/{name}"
        label = "PROFIT" if net_pnl >= 0 else "LOSS"
        name = f"Trade_{tid:04d}_{safe}_{pd.Timestamp(entry_dt).strftime('%Y-%m-%d')}_{label}.png"
        return f"{CHART_REL_PREFIX}/{name}"
    return fn


def sim_statement(
    by_entry: dict,
    paper: dict,
    event_days: list,
    start_dt: pd.Timestamp,
    end_dt: pd.Timestamp,
    ml_n: int,
    closes: CloseCache,
    chart_index: dict,
) -> dict:
    book = simulate_rr_book(
        by_entry, paper, event_days, start_dt, end_dt, closes,
        chart_rel_fn=_year_chart_rel(chart_index),
    )
    tag = f"Start_{start_dt.year}"
    reports = OUT / tag
    plots = PLOTS / tag
    reports.mkdir(parents=True, exist_ok=True)
    plots.mkdir(parents=True, exist_ok=True)

    rows = book["rows"]
    df_daily = book["df_daily"]
    df_stmt = pd.DataFrame(rows)[CSV_COLS]
    df_stmt.to_csv(reports / "Swing_Strategy_Account_Statement.csv", index=False)
    df_daily.to_csv(reports / "Daily_Equity.csv", index=False)

    period = f"{start_dt.strftime('%Y-%m-%d')} to {end_dt.strftime('%Y-%m-%d')}"
    rr_txt = ", ".join(f"{k}={v}" for k, v in sorted(book["rr_fills"].items())) or "none"
    accept_pct = 100.0 * ml_n / max(FEAT_N, 1)
    summary_rows = [
        ("Strategy Rule", RULE, RULE, RULE),
        ("Start vintage", tag, tag, tag),
        ("Initial Capital", f"Rs {CAPITAL:,.2f}", f"Rs {CAPITAL:,.2f}", f"Rs {CAPITAL:,.2f}"),
        ("Risk / Trade", "2% of equity", "2% of equity", "2% of equity"),
        ("Final Portfolio Equity", f"Rs {book['gross_eq']:,.2f}", f"Rs {book['net_z']:,.2f}", f"Rs {book['net_f']:,.2f}"),
        ("Total Net Profit (INR)", f"Rs {book['gross_eq'] - CAPITAL:,.2f}", f"Rs {book['net_z'] - CAPITAL:,.2f}", f"Rs {book['net_f'] - CAPITAL:,.2f}"),
        ("Total Return (%)", f"{book['g_ret']:+.2f}%", f"{book['z_ret']:+.2f}%", f"{book['f_ret']:+.2f}%"),
        ("CAGR (%)", f"{book['g_cagr']:+.2f}%", f"{book['z_cagr']:+.2f}%", f"{book['f_cagr']:+.2f}%"),
        ("Executed Trades Count", f"{book['executed']:,}", f"{book['executed']:,}", f"{book['executed']:,}"),
        ("Win Rate (%)", f"{book['win_rate']:.2f}%", f"{book['win_rate']:.2f}%", f"{book['win_rate']:.2f}%"),
        ("Max Drawdown (%)", f"{book['max_dd']:.2f}%", f"{book['max_dd']:.2f}%", f"{book['max_dd']:.2f}%"),
        ("Total Statutory Taxes Paid", "Rs 0.00", f"Rs {book['total_tax_z']:,.2f}", f"Rs {book['total_tax_f']:,.2f}"),
        ("Backtest Period", period, period, period),
        ("Years (N)", f"{book['years']:.4f}", f"{book['years']:.4f}", f"{book['years']:.4f}"),
        ("Skipped entry days (paper gate)", f"{book['skipped_days']:,}", f"{book['skipped_days']:,}", f"{book['skipped_days']:,}"),
        ("ML accepted / universe", f"{ml_n:,} / {FEAT_N:,}", f"{accept_pct:.1f}%", "Meta_P>=0.48 top16"),
        ("Filled RR mix", rr_txt, rr_txt, rr_txt),
        ("Holding valuation", "MTM qty x close", "MTM qty x close", "MTM qty x close"),
    ]
    _write_excel(rows, summary_rows, reports / "Swing_Strategy_Account_Statement.xlsx")
    generate_all_visualizations(df_daily, plots, reports, exp_title=f"Swing_PP RR {tag}")
    return {
        "start": str(start_dt.date()),
        "end": str(end_dt.date()),
        "years": round(book["years"], 4),
        "Gross_CAGR": round(float(book["g_cagr"]), 2),
        "Zerodha_CAGR": round(float(book["z_cagr"]), 2),
        "FYERS_CAGR": round(float(book["f_cagr"]), 2),
        "Gross_Return": round(float(book["g_ret"]), 2),
        "Zerodha_Return": round(float(book["z_ret"]), 2),
        "Final_Zerodha": round(float(book["net_z"]), 2),
        "Final_Gross": round(float(book["gross_eq"]), 2),
        "DD": round(book["max_dd"], 2),
        "N": book["executed"],
        "Win_Rate": round(book["win_rate"], 2),
        "Taxes_Zerodha": round(book["total_tax_z"], 2),
        "skipped_days": book["skipped_days"],
        "rr_fills": book["rr_fills"],
        "folder": str(reports),
        "open_left": book["open_left"],
        "plot_jobs": book["plot_jobs"],
    }


def _write_master(rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "StartYear CAGR"
    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    headers = [
        "Experiment ID", "Initial Capital", "Risk Cap / Trade", "Start", "End", "Years",
        "Final Account Equity (Zerodha)", "Net Return (%)", "CAGR (%)",
        "Gross CAGR (%)", "FYERS CAGR (%)", "Executed Trades", "Win Rate (%)",
        "Max Drawdown (%)", "Taxes Paid",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in rows:
        ws.append([
            f"Start_{r['start'][:4]}",
            "Rs 50,000",
            "2% of equity",
            r["start"],
            r["end"],
            r["years"],
            f"Rs {r['Final_Zerodha']:,.2f}",
            f"{r['Zerodha_Return']:+.2f}%",
            f"{r['Zerodha_CAGR']:+.2f}%",
            f"{r['Gross_CAGR']:+.2f}%",
            f"{r['FYERS_CAGR']:+.2f}%",
            r["N"],
            f"{r['Win_Rate']:.2f}%",
            f"{r['DD']:.2f}%",
            f"Rs {r['Taxes_Zerodha']:,.2f}",
        ])
    for i, name in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(36, len(name) + 4))
    ws2 = wb.create_sheet("Performance Summary")
    ws2.append(["Performance Metric", "Gross (BEFORE TAX)", "Net (Zerodha)", "Net (Flat Rs 20)"])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws2.append(["See per-vintage Swing_Strategy_Account_Statement.xlsx", "", "", ""])
    wb.save(MASTER)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    paper_src, ml = _scored_fill()
    paper = exits_by_day(paper_src, 1.5)
    by_entry = defaultdict(list)
    for rec in ml.to_dict("records"):
        if pd.isna(rec.get("Exit_Date")):
            continue
        by_entry[pd.Timestamp(rec["Entry_Date"]).normalize()].append(rec)
    for day, cands in by_entry.items():
        cands.sort(key=lambda x: float(x.get("Meta_P", 0.0)), reverse=True)

    end_dt = max(ml["Entry_Date"].max(), pd.to_datetime(ml["Exit_Date"]).max()).normalize()
    extra_x = []
    for recs in by_entry.values():
        for r in recs:
            xd = pd.Timestamp(r["Exit_Date"])
            if pd.notna(xd):
                extra_x.append(xd.normalize())
    event_days = sorted(set(by_entry.keys()) | set(paper.keys()) | set(extra_x))

    closes = CloseCache()
    tickers = {str(r.get("Ticker")) for recs in by_entry.values() for r in recs}
    print(f"Loading {len(tickers)} daily close series for MTM ...", flush=True)
    closes.preload(tickers)
    chart_index = index_trade_charts(CHARTS_2010)
    print(f"trade-chart index {len(chart_index):,} files", flush=True)

    OUT.mkdir(parents=True, exist_ok=True)
    PLOTS.mkdir(parents=True, exist_ok=True)
    years = list(range(2010, end_dt.year + 1))
    print(f"vintages {years[0]}..{years[-1]}  end {end_dt.date()}", flush=True)
    rows = []
    for y in years:
        start_dt = pd.Timestamp(f"{y}-01-01")
        if start_dt > end_dt:
            continue
        rec = sim_statement(by_entry, paper, event_days, start_dt, end_dt, len(ml), closes, chart_index)
        rec.pop("plot_jobs", None)
        rows.append(rec)
        print(
            f"  {rec['start']}  Zerodha CAGR {rec['Zerodha_CAGR']:+7.2f}%  "
            f"DD {rec['DD']:5.1f}%  n={rec['N']:4d}  {rec['years']:.2f}y  "
            f"final Rs {rec['Final_Zerodha']:,.0f}",
            flush=True,
        )

    df = pd.DataFrame(rows)
    df.to_csv(CSV, index=False)
    _write_master(rows)
    LOG.write_text(json.dumps({"end": str(end_dt.date()), "rule": RULE, "rows": rows}, indent=2, default=str), encoding="utf-8")
    print(f"wrote {CSV}", flush=True)
    print(f"wrote {MASTER}", flush=True)


if __name__ == "__main__":
    main()
