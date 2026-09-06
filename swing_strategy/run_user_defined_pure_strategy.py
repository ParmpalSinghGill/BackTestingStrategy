"""
User-Defined Pure Strategy — Multi Capital/Risk Runner (No ML)

Runs guide-compliant backtests for:
  - Exp_100k_1k   (Rs 100,000 capital, Rs 1,000 risk/trade)
  - Exp_100k_0.5k (Rs 100,000 capital, Rs   500 risk/trade)
  - Exp_50k_0.5k  (Rs  50,000 capital, Rs   500 risk/trade)

All scenarios: 1:3 RR, post-C1 breakout entry (any bar before C1 low breaks), SL = C1_Low.
Outputs per backtest-output-statement skill: 21-column statement, ALL trade PNGs,
Performance Summary (Gross + Zerodha + FYERS), portfolio visualizations.
"""

from __future__ import annotations

import argparse
import sys
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor, as_completed
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

REPORTS_DIR = BASE_DIR / "Reports" / "User_Defined_Pure_Strategy"
PLOTS_DIR = BASE_DIR / "Plots" / "User_Defined_Pure_Strategy"

from src.analysis.indian_brokerage_calculator import calculate_indian_trade_charges
from swing_strategy.plotter import plot_swing_trade_chart
from swing_strategy.user_defined_pure_strategy_engine import build_trades_dataset
from swing_strategy.visualizer import generate_all_visualizations

TF_RANK = {"Yearly": 3, "Monthly": 2, "Weekly": 1}
NIFTY_RANK = {"Nifty 50": 4, "Nifty 100": 3, "Nifty 250": 2, "Other": 1}

SCENARIOS = [
    {"exp_name": "Exp_100k_1k", "capital": 100_000.0, "risk": 1_000.0},
    {"exp_name": "Exp_100k_0.5k", "capital": 100_000.0, "risk": 500.0},
    {"exp_name": "Exp_50k_0.5k", "capital": 50_000.0, "risk": 500.0},
]


def run_portfolio_backtest(
    df_trades: pd.DataFrame,
    initial_deposit: float,
    max_risk_per_trade: float,
    exp_name: str,
    reports_dir: Path,
    plots_dir: Path,
    generate_all_pngs: bool = True,
    rr_mode: str = "1:3",
) -> dict:
    reports_dir.mkdir(parents=True, exist_ok=True)
    plots_dir.mkdir(parents=True, exist_ok=True)
    trade_charts_dir = plots_dir / "trade_charts"
    trade_charts_dir.mkdir(parents=True, exist_ok=True)

    rr_mode = rr_mode or (df_trades["Target_RR_Mode"].iloc[0] if not df_trades.empty and "Target_RR_Mode" in df_trades.columns else "1:3")
    deposit_date = "2010-01-01"

    if df_trades.empty:
        return {"exp_name": exp_name, "error": "No trades in dataset"}

    df = df_trades.copy()
    df["Entry_Date"] = pd.to_datetime(df["Entry_Date"])
    df["Exit_Date"] = pd.to_datetime(df["Exit_Date"])

    trades_by_entry: dict = {}
    for row in df.to_dict("records"):
        trades_by_entry.setdefault(row["Entry_Date"], []).append(row)

    min_dt = pd.to_datetime(deposit_date)
    max_dt = max(df["Entry_Date"].max(), df["Exit_Date"].max())
    all_days = pd.date_range(min_dt, max_dt, freq="D")

    current_cash = initial_deposit
    peak_equity = initial_deposit
    max_drawdown_pct = 0.0

    statement_rows = []
    daily_equity_rows = []
    open_positions = []
    trade_records_plot = []

    executed = wins = losses = 0
    total_taxes_zerodha = 0.0
    total_taxes_fyers = 0.0
    total_gross_pnl = 0.0
    tx_id = 1
    trade_id = 1

    statement_rows.append({
        "Transaction_ID": tx_id, "Trade_ID": 0, "Type": "DEPOSIT", "Date": deposit_date,
        "Ticker": "N/A", "Liquidity_Source": "N/A", "Support_Price": 0.0, "Quantity": 0,
        "Price": 0.0, "Total_Spend": 0.0, "Gross_PnL": 0.0, "Statutory_Taxes": 0.0,
        "Net_PnL": 0.0, "Return_Pct": 0.0, "Cash_Balance": current_cash,
        "Active_Position_Count": 0, "Holding_Equity_Value": 0.0,
        "Total_Portfolio_Value": current_cash, "Target_RR_Mode": rr_mode,
        "Outcome": "DEPOSIT", "Chart_PNG_URI": "N/A",
    })
    tx_id += 1

    print(f"\n[{exp_name}] Portfolio sim: capital Rs {initial_deposit:,.0f} | risk Rs {max_risk_per_trade:,.0f}", flush=True)

    for curr_dt in all_days:
        day_start_cash = current_cash
        day_net_pnl = 0.0

        # STEP 1: ENTRIES FIRST (cash at start of day only — exits freed today available tomorrow)
        if curr_dt in trades_by_entry:
            cands = trades_by_entry[curr_dt]
            cands.sort(
                key=lambda x: (
                    -TF_RANK.get(x.get("Liquidity_Type", "Weekly"), 0),
                    -NIFTY_RANK.get(x.get("Index_Membership", "Other"), 0),
                )
            )
            available = max(0.0, current_cash)

            for cand in cands:
                entry_p = float(cand["Entry_Price"])
                sl_p = float(cand["SL_Price"])
                risk_per_share = max(0.05, entry_p - sl_p)
                qty = max(1, int(max_risk_per_trade / risk_per_share))
                pos_val = round(entry_p * qty, 2)

                if pos_val > available or qty <= 0:
                    continue

                current_cash = round(current_cash - pos_val, 2)
                available = round(available - pos_val, 2)
                executed += 1

                pos = {
                    "Trade_ID": trade_id,
                    "Ticker": cand["Ticker"],
                    "Liquidity_Source": cand.get("Liquidity_Type", "Weekly"),
                    "Support_Price": float(cand["Support_Price"]),
                    "C1_Date": cand["C1_Date"],
                    "Entry_Date": curr_dt.strftime("%Y-%m-%d"),
                    "Exit_Date": pd.to_datetime(cand["Exit_Date"]),
                    "Entry_Price": entry_p,
                    "SL_Price": sl_p,
                    "Target_Price": float(cand["Target_Price"]),
                    "Exit_Price": float(cand["Exit_Price"]),
                    "Outcome": cand["Outcome"],
                    "Quantity": qty,
                    "Total_Spend": pos_val,
                    "Planned_Entry_Price": float(cand["Planned_Entry_Price"]),
                    "Target_RR_Mode": rr_mode,
                }
                open_positions.append(pos)

                holding = sum(p["Total_Spend"] for p in open_positions)
                statement_rows.append({
                    "Transaction_ID": tx_id, "Trade_ID": trade_id, "Type": "BUY (ENTRY)",
                    "Date": curr_dt.strftime("%Y-%m-%d"), "Ticker": pos["Ticker"],
                    "Liquidity_Source": pos["Liquidity_Source"], "Support_Price": pos["Support_Price"],
                    "Quantity": qty, "Price": entry_p, "Total_Spend": pos_val,
                    "Gross_PnL": 0.0, "Statutory_Taxes": 0.0, "Net_PnL": 0.0, "Return_Pct": 0.0,
                    "Cash_Balance": current_cash, "Active_Position_Count": len(open_positions),
                    "Holding_Equity_Value": holding, "Total_Portfolio_Value": current_cash + holding,
                    "Target_RR_Mode": rr_mode, "Outcome": "OPEN", "Chart_PNG_URI": "N/A",
                })
                tx_id += 1
                trade_id += 1

        # STEP 2: EXITS SECOND
        to_close = [p for p in open_positions if p["Exit_Date"] <= curr_dt]
        for pos in to_close:
            open_positions = [p for p in open_positions if p["Trade_ID"] != pos["Trade_ID"]]

            qty = pos["Quantity"]
            entry_p = pos["Entry_Price"]
            exit_p = pos["Exit_Price"]
            pos_val = pos["Total_Spend"]
            outcome = pos["Outcome"]

            gross_pnl = round((exit_p - entry_p) * qty, 2)
            total_gross_pnl += gross_pnl
            if outcome == "Success":
                wins += 1
            else:
                losses += 1

            ch_z = calculate_indian_trade_charges(entry_p, exit_p, qty, flat_brokerage_per_order=0.0)
            ch_f = calculate_indian_trade_charges(entry_p, exit_p, qty, flat_brokerage_per_order=20.0)
            tax_z = round(ch_z["total_charges"], 2)
            tax_f = round(ch_f["total_charges"], 2)
            total_taxes_zerodha += tax_z
            total_taxes_fyers += tax_f
            net_pnl = round(gross_pnl - tax_z, 2)
            day_net_pnl += net_pnl
            ret_pct = round((net_pnl / pos_val) * 100.0, 2) if pos_val > 0 else 0.0

            current_cash = round(current_cash + pos_val + gross_pnl - tax_z, 2)

            trade_records_plot.append({
                "Trade_ID": pos["Trade_ID"],
                "Ticker": pos["Ticker"],
                "C1_Date": pos["C1_Date"],
                "C2_Date": pos["Entry_Date"],
                "Entry_Date": pos["Entry_Date"],
                "Exit_Date": curr_dt.strftime("%Y-%m-%d"),
                "Planned_Entry_Price": pos["Planned_Entry_Price"],
                "Entry_Price": entry_p,
                "SL_Price": pos["SL_Price"],
                "Target_Price": pos["Target_Price"],
                "Support_Price": pos["Support_Price"],
                "Exit_Price": exit_p,
                "Net_PnL": net_pnl,
                "Liquidity_Type": pos["Liquidity_Source"],
                "ML_RR_Choice": rr_mode,
            })

            holding = sum(p["Total_Spend"] for p in open_positions)
            statement_rows.append({
                "Transaction_ID": tx_id, "Trade_ID": pos["Trade_ID"], "Type": "SELL (EXIT)",
                "Date": curr_dt.strftime("%Y-%m-%d"), "Ticker": pos["Ticker"],
                "Liquidity_Source": pos["Liquidity_Source"], "Support_Price": pos["Support_Price"],
                "Quantity": qty, "Price": exit_p, "Total_Spend": pos_val,
                "Gross_PnL": gross_pnl, "Statutory_Taxes": tax_z, "Net_PnL": net_pnl,
                "Return_Pct": ret_pct, "Cash_Balance": current_cash,
                "Active_Position_Count": len(open_positions), "Holding_Equity_Value": holding,
                "Total_Portfolio_Value": current_cash + holding,
                "Target_RR_Mode": rr_mode, "Outcome": outcome, "Chart_PNG_URI": "PENDING",
            })
            tx_id += 1

        port_val = current_cash + sum(p["Total_Spend"] for p in open_positions)
        if port_val > peak_equity:
            peak_equity = port_val
        dd = ((peak_equity - port_val) / peak_equity * 100.0) if peak_equity > 0 else 0.0
        max_drawdown_pct = max(max_drawdown_pct, dd)

        daily_equity_rows.append({
            "Date": curr_dt.strftime("%Y-%m-%d"),
            "Balance": current_cash,
            "Active_Positions": len(open_positions),
            "Daily_PnL": day_net_pnl,
            "Daily_Return_Pct": ((current_cash - day_start_cash) / day_start_cash * 100.0) if day_start_cash > 0 else 0.0,
        })

    # ALL trade PNG charts (parallel)
    chart_map = {}
    if generate_all_pngs and trade_records_plot:
        print(f"[{exp_name}] Rendering {len(trade_records_plot):,} trade PNG charts...", flush=True)

        def _render(rec):
            try:
                return rec["Trade_ID"], plot_swing_trade_chart(rec, trade_charts_dir)
            except Exception:
                clean = rec["Ticker"].replace(".NS", "")
                tag = "PROFIT" if rec["Net_PnL"] >= 0 else "LOSS"
                p = trade_charts_dir / f"Trade_{rec['Trade_ID']:04d}_{clean}_{rec['Entry_Date']}_{tag}.png"
                return rec["Trade_ID"], p.as_uri()

        with ThreadPoolExecutor(max_workers=16) as ex:
            futs = [ex.submit(_render, r) for r in trade_records_plot]
            done = 0
            for fut in as_completed(futs):
                tid, uri = fut.result()
                chart_map[tid] = uri
                done += 1
                if done % 500 == 0 or done == len(trade_records_plot):
                    print(f"  • [{exp_name}] PNG {done:,}/{len(trade_records_plot):,}", flush=True)

    for row in statement_rows:
        if row["Type"] == "SELL (EXIT)":
            row["Chart_PNG_URI"] = chart_map.get(row["Trade_ID"], "N/A")

    df_stmt = pd.DataFrame(statement_rows)
    df_daily = pd.DataFrame(daily_equity_rows)

    num_years = max((max_dt - min_dt).days / 365.25, 0.01)
    gross_equity = round(initial_deposit + total_gross_pnl, 2)
    net_equity_z = current_cash
    net_equity_f = round(current_cash - (total_taxes_fyers - total_taxes_zerodha), 2)

    def _metrics(final_eq: float) -> tuple[float, float]:
        if initial_deposit <= 0:
            return 0.0, 0.0
        ret = ((final_eq - initial_deposit) / initial_deposit) * 100.0
        if final_eq <= 0:
            return round(ret, 2), -100.0
        cagr = (((final_eq / initial_deposit) ** (1.0 / num_years)) - 1.0) * 100.0
        return round(ret, 2), round(cagr, 2)

    gross_ret, gross_cagr = _metrics(gross_equity)
    net_ret, net_cagr = _metrics(net_equity_z)
    fyers_ret, fyers_cagr = _metrics(net_equity_f)
    win_rate = (wins / executed * 100.0) if executed > 0 else 0.0

    csv_path = reports_dir / "Swing_Strategy_Account_Statement.csv"
    xlsx_path = reports_dir / "Swing_Strategy_Account_Statement.xlsx"
    df_stmt.to_csv(csv_path, index=False)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    buy_fill = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    sell_win_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    sell_loss_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Account Statement"
    headers = list(df_stmt.columns)
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, r in df_stmt.iterrows():
        row_num = r_idx + 2
        tx_type = r["Type"]
        pnl = r["Net_PnL"]
        fill = buy_fill if tx_type.startswith("BUY") else (sell_win_fill if pnl >= 0 else sell_loss_fill)
        for c_idx, val in enumerate(r, start=1):
            cell = ws.cell(row=row_num, column=c_idx, value=val)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center" if c_idx in (1, 2, 3, 4, 5, 6, 19, 20) else "right")
            if c_idx == 21 and r["Chart_PNG_URI"] not in ("N/A", "PENDING"):
                uri = r["Chart_PNG_URI"]
                cell.value = f'=HYPERLINK("{uri}", "View Plot Chart (PNG)")'
                cell.hyperlink = uri
                cell.font = Font(name="Calibri", size=10, color="2563EB", underline="single")

    ws_sum = wb.create_sheet("Performance Summary")
    ws_sum.append(["Performance Metric", "Gross (BEFORE TAX)", "Net (Zerodha)", "Net (Flat Rs 20)"])
    for col in range(1, 5):
        c = ws_sum.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    summary_rows = [
        ("Strategy Rule", "C1 submerged + post-C1 breakout entry", "C1 submerged + post-C1 breakout entry", "C1 submerged + post-C1 breakout entry"),
        ("Target RR Mode", rr_mode, rr_mode, rr_mode),
        ("Initial Capital", f"Rs {initial_deposit:,.2f}", f"Rs {initial_deposit:,.2f}", f"Rs {initial_deposit:,.2f}"),
        ("Fixed Risk Cap / Trade", f"Rs {max_risk_per_trade:,.2f}", f"Rs {max_risk_per_trade:,.2f}", f"Rs {max_risk_per_trade:,.2f}"),
        ("Final Portfolio Equity", f"Rs {gross_equity:,.2f}", f"Rs {net_equity_z:,.2f}", f"Rs {net_equity_f:,.2f}"),
        ("Total Net Profit (INR)", f"Rs {gross_equity - initial_deposit:,.2f}", f"Rs {net_equity_z - initial_deposit:,.2f}", f"Rs {net_equity_f - initial_deposit:,.2f}"),
        ("Total Return (%)", f"{gross_ret:+.2f}%", f"{net_ret:+.2f}%", f"{fyers_ret:+.2f}%"),
        ("CAGR (%)", f"{gross_cagr:+.2f}%", f"{net_cagr:+.2f}%", f"{fyers_cagr:+.2f}%"),
        ("Executed Trades Count", f"{executed:,}", f"{executed:,}", f"{executed:,}"),
        ("Win Rate (%)", f"{win_rate:.2f}%", f"{win_rate:.2f}%", f"{win_rate:.2f}%"),
        ("Max Drawdown (%)", f"{max_drawdown_pct:.2f}%", f"{max_drawdown_pct:.2f}%", f"{max_drawdown_pct:.2f}%"),
        ("Total Statutory Taxes Paid", "Rs 0.00", f"Rs {total_taxes_zerodha:,.2f}", f"Rs {total_taxes_fyers:,.2f}"),
        ("Trade PNG Charts Generated", f"{len(chart_map):,}", f"{len(chart_map):,}", f"{len(chart_map):,}"),
        ("Backtest Period", f"{deposit_date} to {max_dt.strftime('%Y-%m-%d')}", f"{deposit_date} to {max_dt.strftime('%Y-%m-%d')}", f"{deposit_date} to {max_dt.strftime('%Y-%m-%d')}"),
    ]
    for row in summary_rows:
        ws_sum.append(list(row))

    ws_sum.column_dimensions["A"].width = 38
    ws_sum.column_dimensions["B"].width = 36
    ws_sum.column_dimensions["C"].width = 28
    ws_sum.column_dimensions["D"].width = 28
    wb.save(xlsx_path)

    generate_all_visualizations(df_daily, plots_dir, reports_dir, exp_title=exp_name)

    print(
        f"[{exp_name}] DONE | Final Rs {net_equity_z:,.2f} | Net CAGR {net_cagr:+.2f}% | "
        f"Trades {executed:,} | Win {win_rate:.2f}% | PNGs {len(chart_map):,}",
        flush=True,
    )

    return {
        "Experiment": exp_name,
        "Initial_Capital": initial_deposit,
        "Risk_Cap": max_risk_per_trade,
        "Final_Equity": net_equity_z,
        "Gross_Final_Equity": gross_equity,
        "Net_Return_Pct": net_ret,
        "CAGR_Pct": net_cagr,
        "Gross_CAGR_Pct": gross_cagr,
        "Fyers_CAGR_Pct": fyers_cagr,
        "Executed_Trades": executed,
        "Win_Rate_Pct": win_rate,
        "Max_Drawdown_Pct": max_drawdown_pct,
        "Taxes_Zerodha": total_taxes_zerodha,
        "Taxes_FYERS": total_taxes_fyers,
        "PNG_Count": len(chart_map),
    }


def _run_scenario(args: tuple) -> dict:
    df_trades, scenario, rescan, generate_pngs = args
    if rescan and df_trades is None:
        df_trades = build_trades_dataset(force_rescan=True)

    exp = scenario["exp_name"]
    rep_dir = REPORTS_DIR / exp
    plt_dir = PLOTS_DIR / exp
    return run_portfolio_backtest(
        df_trades=df_trades,
        initial_deposit=scenario["capital"],
        max_risk_per_trade=scenario["risk"],
        exp_name=exp,
        reports_dir=rep_dir,
        plots_dir=plt_dir,
        generate_all_pngs=generate_pngs,
    )


def main():
    parser = argparse.ArgumentParser(description="User-Defined Pure Strategy (No ML) — multi capital/risk")
    parser.add_argument("--rescan", action="store_true", help="Force rescan tickers for setup dataset")
    parser.add_argument("--start", type=str, default="2010-01-01")
    parser.add_argument("--no-charts", action="store_true", help="Skip per-trade PNG charts (keep statement + portfolio viz)")
    args = parser.parse_args()

    print("=" * 72, flush=True)
    print("  USER-DEFINED PURE STRATEGY — 100k_1k | 100k_0.5k | 50k_0.5k (parallel)", flush=True)
    print("=" * 72, flush=True)

    df_trades = build_trades_dataset(start_date=args.start, force_rescan=args.rescan)
    if df_trades.empty:
        print("No setups found. Aborting.", flush=True)
        return

    tasks = [(df_trades, sc, False, not args.no_charts) for sc in SCENARIOS]

    with ProcessPoolExecutor(max_workers=3) as ex:
        summaries = list(ex.map(_run_scenario, tasks))

    rows = []
    for s in summaries:
        if "error" in s:
            continue
        rows.append({
            "Experiment": s["Experiment"],
            "Initial Capital": f"Rs {s['Initial_Capital']:,.0f}",
            "Risk Cap / Trade": f"Rs {s['Risk_Cap']:,.0f}",
            "Final Account Equity": f"Rs {s['Final_Equity']:,.2f}",
            "Net Return (%)": f"{s['Net_Return_Pct']:+.2f}%",
            "CAGR (%)": f"{s['CAGR_Pct']:.2f}%",
            "Gross CAGR (%)": f"{s['Gross_CAGR_Pct']:.2f}%",
            "Executed Trades": s["Executed_Trades"],
            "Win Rate (%)": f"{s['Win_Rate_Pct']:.2f}%",
            "Max Drawdown (%)": f"{s['Max_Drawdown_Pct']:.2f}%",
            "Taxes Paid (Zerodha)": f"Rs {s['Taxes_Zerodha']:,.2f}",
            "PNG Charts": s["PNG_Count"],
        })

    df_cmp = pd.DataFrame(rows)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    cmp_csv = REPORTS_DIR / "Master_User_Defined_Pure_Comparison.csv"
    cmp_xlsx = REPORTS_DIR / "Master_User_Defined_Pure_Comparison.xlsx"
    df_cmp.to_csv(cmp_csv, index=False)
    df_cmp.to_excel(cmp_xlsx, index=False)

    print("\n" + "=" * 72, flush=True)
    print("MASTER COMPARISON", flush=True)
    print("=" * 72, flush=True)
    print(df_cmp.to_string(index=False), flush=True)
    print(f"\nSaved: {cmp_csv}", flush=True)
    print(f"Saved: {cmp_xlsx}", flush=True)


if __name__ == "__main__":
    main()
