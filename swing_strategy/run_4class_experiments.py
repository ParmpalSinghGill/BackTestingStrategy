"""
Parallel Experiment Runner: Single 4-Class Model vs Two-Stage Cascade ML Model

Evaluates under 100k Capital & 1k Risk Cap (Following Guide/Realistic_guide.md):
1. Single 4-Class Model: Multi-class classifier predicting Skip, 1:2, 1:3, or 1:4.
2. Two-Stage Cascade Model: Stage 1 (Trade vs Skip) + Stage 2 (1:2 vs 1:3 vs 1:4 Target Selector).

Executes both models simultaneously in parallel processes for ultra-fast performance comparison.
Exports Master_4Class_Comparison.xlsx side-by-side comparative report.
"""

import os
import sys
from pathlib import Path
import pandas as pd
from concurrent.futures import ProcessPoolExecutor

# Force UTF-8 encoding for stdout on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

COMPARISON_DIR = BASE_DIR / "Reports" / "Model_Comparison_4Class_vs_2Stage"
PLOTS_DIR = BASE_DIR / "Plots" / "Model_Comparison_4Class_vs_2Stage"

os.makedirs(COMPARISON_DIR, exist_ok=True)
os.makedirs(PLOTS_DIR, exist_ok=True)

from swing_strategy.strategy_4class_engine import (
    prepare_4class_dataset,
    run_single_4class_model,
    run_two_stage_cascade_model
)
from swing_strategy.generate_statement import generate_swing_strategy_statement


def run_single_4class_task(df_4c):
    exp_name = "Single_4Class_Model"
    rep_dir = COMPARISON_DIR / exp_name
    plt_dir = PLOTS_DIR / exp_name

    df_acc_4class = run_single_4class_model(df_4c, probability_threshold=0.40)

    summary = generate_swing_strategy_statement(
        initial_deposit=100000.0,
        max_charts_to_generate=25,
        max_risk_per_trade=1000.0,
        df_acc=df_acc_4class,
        custom_reports_dir=rep_dir,
        custom_plots_dir=plt_dir,
        exp_name=exp_name
    )
    return summary


def run_two_stage_cascade_task(df_4c):
    exp_name = "Two_Stage_Cascade_Model"
    rep_dir = COMPARISON_DIR / exp_name
    plt_dir = PLOTS_DIR / exp_name

    df_acc_2stage = run_two_stage_cascade_model(df_4c, stage1_threshold=0.42)

    summary = generate_swing_strategy_statement(
        initial_deposit=100000.0,
        max_charts_to_generate=25,
        max_risk_per_trade=1000.0,
        df_acc=df_acc_2stage,
        custom_reports_dir=rep_dir,
        custom_plots_dir=plt_dir,
        exp_name=exp_name
    )
    return summary


def run_comparison():
    print("==========================================================================", flush=True)
    print("     4-CLASS SINGLE MODEL VS 2-STAGE CASCADE ML MODEL COMPARISON SUITE    ", flush=True)
    print("==========================================================================", flush=True)

    # Prepare dataset once
    df_4c = prepare_4class_dataset()

    print("\n--- Running Both Models SIMULTANEOUSLY in Parallel Processes ---\n", flush=True)

    with ProcessPoolExecutor(max_workers=2) as executor:
        f1 = executor.submit(run_single_4class_task, df_4c)
        f2 = executor.submit(run_two_stage_cascade_task, df_4c)

        s1 = f1.result()
        s2 = f2.result()

    summary_list = [s1, s2]

    display_rows = []
    for s in summary_list:
        display_rows.append({
            "ML Architecture": s["Exp_Name"],
            "Initial Deposit": f"Rs {s['Initial_Capital']:,.0f}",
            "Risk Cap": f"Rs {s['Max_Risk_Cap']:,.0f}",
            "Final Account Equity": f"Rs {s['Final_Equity']:,.2f}",
            "Net Return (%)": f"+{s['Total_Net_Return_Pct']:,.2f}%",
            "CAGR (%)": f"{s['CAGR_Pct']:.2f}%",
            "Executed Trades": f"{s['Executed_Trades']:,}",
            "Win Rate (%)": f"{s['Win_Rate_Pct']:.2f}%",
            "Max Drawdown (%)": f"{s['Max_Drawdown_Pct']:.2f}%",
            "Target Distribution (1:2 / 1:3 / 1:4)": f"{s['Count_1to2']:,} / {s['Count_1to3']:,} / {s['Count_1to4']:,}",
            "Taxes Paid": f"Rs {s['Total_Taxes_Paid']:,.2f}"
        })
    df_display = pd.DataFrame(display_rows)

    print("\n==========================================================================================", flush=True)
    print("                    4-CLASS VS 2-STAGE CASCADE MODEL PERFORMANCE COMPARISON                ", flush=True)
    print("==========================================================================================", flush=True)
    print(df_display.to_string(index=False), flush=True)
    print("==========================================================================================\n", flush=True)

    master_excel_path = COMPARISON_DIR / "Master_4Class_Comparison.xlsx"
    master_csv_path = COMPARISON_DIR / "Master_4Class_Comparison.csv"

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "4-Class Models Comparison"

        headers = list(df_display.columns)
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, row in df_display.iterrows():
            row_data = list(row)
            ws.append(row_data)

        wb.save(master_excel_path)
        print(f"Master Comparison Excel saved to: {master_excel_path.resolve()}", flush=True)
    except Exception as e:
        print(f"Excel export warning: {e}", flush=True)

    df_display.to_csv(master_csv_path, index=False)
    return master_excel_path, df_display


if __name__ == "__main__":
    run_comparison()
